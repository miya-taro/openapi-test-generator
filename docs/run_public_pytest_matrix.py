#!/usr/bin/env python3
"""Opt-in, low-volume public API checks through the application's native pytest runner.

No public requests on import or preparation. One case per subprocess prevents
parallel bursts; stop a host on access/rate/transport failures, without retries.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import time
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'runner'))

from docs.run_public_api_checks import specifications, selected_cases, BASE_URLS, obj
from docs.run_public_rest_checks import profiles, HOSTS, ECHO
from docs.run_public_generated_checks import prepare, BASE_URL
from workflow.case_master import build_case_master, write_case_master
from openpyxl import Workbook


def additional_profiles():
    text, number, integer = {'type': 'string'}, {'type': 'number'}, {'type': 'integer'}
    definitions = [
        ('openmeteo', 'https://api.open-meteo.com', '/v1/forecast', 'GET',
         obj({'latitude': number, 'longitude': number, 'current': obj({'temperature_2m': number, 'time': text})}),
         {'latitude': 35.68, 'longitude': 139.69, 'current': 'temperature_2m'}, None,
         'https://open-meteo.com/en/docs'),
        ('pokeapi', 'https://pokeapi.co', '/api/v2/pokemon/25/', 'GET',
         obj({'id': {'const': 25}, 'name': {'const': 'pikachu'}, 'stats': {'type': 'array', 'minItems': 1,
              'items': obj({'base_stat': integer, 'stat': obj({'name': text, 'url': {'type': 'string', 'format': 'uri'}})})}}),
         {}, None, 'https://pokeapi.co/docs/v2#pokemon'),
        ('restcountries', 'https://restcountries.com', '/v3.1/alpha/jp', 'GET',
         {'type': 'array', 'minItems': 1, 'maxItems': 1, 'items': obj({'cca2': {'const': 'JP'}, 'name': obj({'common': {'const': 'Japan'}})})},
         {}, None, 'https://restcountries.com/'),
        ('worldbank', 'https://api.worldbank.org', '/v2/country/JPN/indicator/SP.POP.TOTL', 'GET',
         {'type': 'array', 'minItems': 2, 'maxItems': 2, 'prefixItems': [
             obj({'page': integer, 'pages': integer, 'total': integer}),
             {'type': 'array', 'minItems': 1, 'items': obj({'countryiso3code': {'const': 'JPN'},
              'date': {'const': '2020'}, 'value': {'type': ['number', 'null']}})}]},
         {'format': 'json', 'date': '2020'}, None,
         'https://datahelpdesk.worldbank.org/knowledgebase/articles/898581-api-basic-call-structures'),
        ('goproxy', 'https://proxy.golang.org', '/github.com/gin-gonic/gin/@v/v1.10.0.info', 'GET',
         obj({'Version': {'const': 'v1.10.0'}, 'Time': {'type': 'string', 'format': 'date-time'}}),
         {}, None, 'https://go.dev/ref/mod#goproxy-protocol'),
        ('rubygems', 'https://rubygems.org', '/api/v1/gems/rails.json', 'GET',
         obj({'name': {'const': 'rails'}, 'version': text, 'downloads': integer,
              'dependencies': obj({'development': {'type': 'array'}, 'runtime': {'type': 'array'}})}),
         {}, None, 'https://guides.rubygems.org/rubygems-org-api/'),
        ('nuget', 'https://api.nuget.org', '/v3/registration5-semver1/newtonsoft.json/index.json', 'GET',
         obj({'count': {'type': 'integer', 'minimum': 1}, 'items': {'type': 'array', 'minItems': 1,
              'items': obj({'count': integer, 'lower': text, 'upper': text})}}),
         {}, None, 'https://learn.microsoft.com/en-us/nuget/api/registration-base-url-resource'),
        ('postmanecho', 'https://postman-echo.com', '/post', 'POST',
         obj({'json': {'const': ECHO}}), {}, ECHO, 'https://learning.postman.com/docs/developer/echo-api/'),
    ]
    for name, base, path, method, schema, params, body, source in definitions:
        op = {'operationId': name, 'responses': {'200': {'description': 'Selected documented response fields',
              'content': {'application/json': {'schema': schema}}}},
              'parameters': [{'name': key, 'in': 'query', 'required': True,
                              'schema': {'type': 'number' if isinstance(value, (float, int)) else 'string', 'example': value}}
                             for key, value in params.items()]}
        # Required here fixes the test invocation, not the service's optionality contract.
        if body is not None:
            op['requestBody'] = {'required': True, 'content': {'application/json': {'schema': {'type': 'object', 'example': body}}}}
        spec = {'openapi': '3.1.0', 'info': {'title': name + ' selected checks', 'version': '1', 'description': source},
                'servers': [{'url': base}], 'paths': {path: {method.lower(): op}}}
        case = {'id': 'PUBLIC-' + name, 'operationId': name, 'summary': name + ' documented example',
                'path': path, 'method': method, 'in': 'body' if body else 'operation',
                'paramName': 'requestBody' if body else '__valid_request__',
                'inputValue': json.dumps(body, ensure_ascii=False) if body else '', 'expectedStatus': '200',
                'perspective': '正常系', 'keyword': 'explicit-public-contract'}
        yield name, base, spec, [case], 'authored-example'


def matrix():
    specs, cases = profiles()
    for name in specs:
        yield name, HOSTS[name], specs[name], cases[name], 'authored-compatibility'
    for name, spec in specifications().items():
        if name in {'pypi', 'npm'}:
            # Explicit compatibility snapshots from 2026-10-20, NOT independent
            # business or official error-body specifications. Keep status=404.
            baseline = {'message': 'Not Found'} if name == 'pypi' else 'Not Found'
            spec['info']['description'] = '404 body: observed compatibility snapshot on 2026-10-20; not an independently specified error message'
            for path_item in spec['paths'].values():
                path_item['get']['responses']['404']['content'] = {'application/json': {'schema': {'const': baseline}}}
        yield name, BASE_URLS[name], spec, selected_cases(spec, name), 'generated-baseline-and-authored-variants'
    spec, cases = prepare()
    yield 'crates-generated', BASE_URL, spec, cases, 'all-generated-cases'
    yield from additional_profiles()


def save_summary(output, summary):
    (output / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    book = Workbook(); sheet = book.active; sheet.title = '公開API検証'
    sheet.append(['サービス', '試験', 'HTTPメソッド', '期待status', '実測status', '判定', '実行状態', '詳細'])
    for row in summary['cases']:
        sheet.append([row.get(key, '') for key in ('provider', 'title', 'method', 'expectedStatus', 'actualStatus', 'verdict', 'status', 'detail')])
    sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip('ABCDEFGH', [22, 55, 16, 16, 16, 12, 18, 80], strict=True): sheet.column_dimensions[column].width = width
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in sheet[1]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='24476B')
    for row in sheet.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(vertical='top', wrap_text=True)
    book.save(output / 'summary.xlsx'); book.close()


def audit_responses(output, summary):
    """Offline fault injection into recorded public responses; zero HTTP sends."""
    from dataclasses import replace
    from execution_plan import contract_case
    from models import RunResult
    from response_schema import from_run_case
    from verifier import verify
    controls = []
    for entry in summary['cases']:
        if entry['status'] != 'passed':
            continue
        master = json.loads((output / entry['provider'] / 'case-master.json').read_text(encoding='utf-8'))
        case = next(case for case in master['cases'] if case['caseId'] == entry['caseId'])
        api = master['execution']['openapi']
        record = json.loads((output / entry['evidence']).read_text(encoding='utf-8'))['cases'][0]
        current = contract_case(case)
        result = RunResult(current, record['actualStatus'], record['body'], record['headers'], record['elapsedMs'], 0)
        contract = from_run_case(current, api)
        assert verify(result, contract)[0] == 'OK', entry['caseId']
        variants = [('unexpected_500', replace(result, actual_status=500))]
        if contract.expected_body_schemas_by_status.get(record['actualStatus']) is not None:
            variants.append(('invalid_json_body', replace(result, actual_body='not valid JSON', report_body=None)))
        if record['actualStatus'] in contract.expected_empty_body_statuses:
            variants.append(('unexpected_response_body', replace(result, actual_body='unexpected body', report_body=None)))
        for name, altered in variants:
            verdict, detail = verify(altered, contract)
            controls.append({'caseId': entry['caseId'], 'provider': entry['provider'], 'mutation': name,
                             'expected': 'NG', 'actual': verdict, 'detail': detail})
    summary['offlineControls'] = controls
    summary['offlineControlsPassed'] = sum(row['expected'] == row['actual'] for row in controls)
    return all(row['expected'] == row['actual'] for row in controls)


def run(output, execute=False, providers=None, case_ids=None):
    output.mkdir(parents=True, exist_ok=True)
    summary = {'mode': 'executed' if execute else 'prepared-only', 'engine': 'native-pytest',
               'startedAt': datetime.now(timezone.utc).isoformat(), 'cases': [], 'groups': []}
    summary['gitCommit'] = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip()
    summary['scriptSha256'] = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    (output / 'execution-script.py').write_bytes(Path(__file__).read_bytes())
    env = os.environ.copy()
    for key in list(env):
        if key.startswith('TEST_AUTH_'): env.pop(key)
    env['NETRC'] = os.devnull  # No unrelated login credentials; retain proxy and CA environment.
    stopped_hosts = {}
    for provider, base, spec, cases, origin in matrix():
        if providers and provider not in providers: continue
        if case_ids:
            cases = [case for case in cases if case['id'] in case_ids]
            if not cases: continue
        for path_item in spec['paths'].values():
            for method, op in path_item.items():
                if method not in {'get', 'post', 'put', 'patch', 'delete'}: continue
                op.setdefault('parameters', []).append({'name': 'User-Agent', 'in': 'header', 'required': True,
                    'schema': {'type': 'string', 'example': 'openapi-test-generator/0.1 (https://github.com/miya-taro/openapi-test-generator)'}})
        folder = output / provider; folder.mkdir(exist_ok=True)
        spec_path = folder / 'openapi.json'
        spec_path.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding='utf-8')
        for case in cases: case['expectedResponseTime'] = ''
        master = build_case_master(provider, {'cases': cases}, None, spec_path)
        master_path = folder / 'case-master.json'; write_case_master(master, master_path)
        summary['groups'].append({'provider': provider, 'baseUrl': base, 'origin': origin, 'planned': len(cases),
                                  'specSha256': hashlib.sha256(spec_path.read_bytes()).hexdigest()})
        for index, case in enumerate(master['cases'], 1):
            entry = {'provider': provider, 'caseId': case['caseId'], 'title': case['title'],
                     'method': case['definition']['request']['method'], 'expectedStatus': case['expected']['status'],
                     'status': 'not-run', 'verdict': '未実施', 'actualStatus': 0}
            summary['cases'].append(entry)
            if not execute or base in stopped_hosts:
                entry['detail'] = stopped_hosts.get(base, 'Prepared without sending HTTP')
                continue
            destination = folder / f'{index:02d}'; destination.mkdir()
            time.sleep(1)
            command = [sys.executable, str(ROOT / 'runner/pytest_runner.py'), str(master_path), '--kind', 'contract',
                       '--base-url', base, '--skip-health-check', '--case-ids', case['caseId'],
                       '--output', str(destination / 'result.xlsx'), '--junit-xml', str(destination / 'junit.xml'), '--full-body']
            with (destination / 'run.log').open('w', encoding='utf-8') as log:
                try:
                    completed = subprocess.run(command, env=env, cwd=ROOT, stdout=log, stderr=subprocess.STDOUT, timeout=90, check=False)
                    entry['processExitCode'] = completed.returncode
                except subprocess.TimeoutExpired:
                    stopped_hosts[base] = 'Runner timed out; subsequent requests stopped'
                    entry.update(status='error', verdict='NG', detail=stopped_hosts[base]); continue
            result_path = destination / 'result.json'
            if not result_path.exists():
                entry.update(status='error', verdict='NG', detail='Runner produced no result.json; inspect run.log')
                stopped_hosts[base] = entry['detail']; continue
            result = json.loads(result_path.read_text(encoding='utf-8'))
            record = result['cases'][0]
            for key in ('status', 'verdict', 'actualStatus', 'detail', 'elapsedMs', 'executedAt', 'curl', 'error'):
                if key in record: entry[key] = record[key]
            entry['evidence'] = str(result_path.relative_to(output))
            # Native JUnit must represent precisely the same one test and verdict.
            junit = ET.parse(destination / 'junit.xml').getroot()
            items = list(junit.iter('testcase'))
            assert len(items) == 1 and items[0].get('name') == case['caseId']
            assert (items[0].find('failure') is None and items[0].find('error') is None and items[0].find('skipped') is None) == (entry['status'] == 'passed')
            if record.get('error') or record.get('actualStatus') in (403, 429) or record.get('actualStatus', 0) >= 500:
                stopped_hosts[base] = 'Transport/access/rate/server error; remaining cases not run'
            print(provider, case['caseId'], entry['actualStatus'], entry['status'], entry.get('detail', '')[:180], flush=True)
            save_summary(output, summary)
    summary['finishedAt'] = datetime.now(timezone.utc).isoformat()
    summary['counts'] = dict(Counter(row['status'] for row in summary['cases']))
    if execute: audit_responses(output, summary)
    save_summary(output, summary)
    print(json.dumps(summary['counts']), flush=True)
    return summary


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--execute', action='store_true')
    parser.add_argument('--providers', nargs='+', choices=['jsonplaceholder', 'httpbin', 'pypi', 'npm', 'crates', 'crates-generated', 'openmeteo', 'pokeapi', 'restcountries', 'worldbank', 'postmanecho', 'goproxy', 'rubygems', 'nuget'], help='Run only these provider groups')
    parser.add_argument('--case-ids', nargs='+', help='Run only these case IDs')
    parser.add_argument('--output-dir', type=Path, required=True, help='New directory for this run (no overwriting evidence)')
    args = parser.parse_args()
    destination = args.output_dir.resolve()
    if destination.exists(): parser.error('Choose a new output directory')
    result = run(destination, args.execute, args.providers, args.case_ids)
    raise SystemExit(1 if not result['cases'] or (args.execute and (any(row['status'] != 'passed' for row in result['cases']) or result.get('offlineControlsPassed') != len(result.get('offlineControls', [])))) else 0)
