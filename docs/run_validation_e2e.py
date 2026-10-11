#!/usr/bin/env python3
"""Generate and execute request/response contract cases against validation_api.py.

Run from the repository root:
    python docs/run_validation_e2e.py

Pass --output-dir to retain the generated JSON, Excel specification, and result
workbook. The local API and generator/runner communicate over HTTP on loopback;
no external host, database, or credentials are used.
"""
from __future__ import annotations

import argparse
import json
import re
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SPEC = ROOT / "docs" / "validation_openapi.json"


def _free_port() -> int:
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        return int(listener.getsockname()[1])


def _wait_for_health(base_url: str, process: subprocess.Popen, timeout: float = 15.0) -> None:
    deadline = time.monotonic() + timeout
    last_error: Exception | None = None
    while time.monotonic() < deadline:
        if process.poll() is not None:
            raise RuntimeError(f"validation API exited during startup (code {process.returncode})")
        try:
            with urllib.request.urlopen(f"{base_url}/health", timeout=1) as response:
                if response.status == 200:
                    return
        except (urllib.error.URLError, OSError) as error:
            last_error = error
        time.sleep(0.1)
    raise RuntimeError(f"validation API did not become healthy: {last_error}")


def _run(command: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(command, cwd=ROOT, text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)


def _check_combined_required_inputs(cases: list[dict], result_path: Path) -> list[str]:
    """Ensure record cases keep all other required inputs valid around each target."""
    workbook = load_workbook(result_path, data_only=True, read_only=True)
    try:
        sheet = workbook["test_spec"]
        columns = {cell.value: index for index, cell in enumerate(sheet[4]) if cell.value}
        rows_by_id = {
            str(row[columns["試験ID"]]): row
            for row in sheet.iter_rows(min_row=5, values_only=True)
            if row[columns["試験ID"]]
        }
        required_inputs = [
            ("query", "tenant", "tenant=north"),
            ("header", "X-Request-Source", "X-Request-Source: contract-test"),
            ("cookie", "session_id", "session_id=sess-ab12"),
        ]
        errors: list[str] = []
        for case in cases:
            if case["operationId"] != "createValidationRecord":
                continue
            row = rows_by_id.get(case["id"])
            command = str(row[columns["実行コマンド"]] or "") if row else ""
            if not command:
                errors.append(f"{case['id']}: executed request command is missing")
                continue

            target = (case["in"], case["paramName"])
            for location, name, expected in required_inputs:
                if target == (location, name):
                    continue
                if expected not in command:
                    errors.append(f"{case['id']}: required {location} parameter {name} was not sent")

            omitted_body = (
                case["in"] == "body"
                and case["paramName"] == "requestBody"
                and case["inputValue"] in {"（省略）", "（フィールドを省略）"}
            )
            has_body = " -d " in command
            if has_body == omitted_body:
                errors.append(f"{case['id']}: required request body was not isolated to its target case")
        return errors
    finally:
        workbook.close()


def _execute(output_dir: Path) -> int:
    output_dir.mkdir(parents=True, exist_ok=True)
    port = _free_port()
    base_url = f"http://127.0.0.1:{port}"
    api = subprocess.Popen(
        [
            sys.executable, "-m", "uvicorn", "docs.validation_api:app",
            "--host", "127.0.0.1", "--port", str(port), "--log-level", "error",
        ],
        cwd=ROOT,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    try:
        _wait_for_health(base_url, api)
        cases_path = output_dir / "cases.json"
        spec_path = output_dir / "test_spec.xlsx"
        result_path = output_dir / "result.xlsx"

        generated = _run([
            sys.executable, "src/cli.py", str(SPEC),
            "--output", str(cases_path), "--excel-output", str(spec_path),
        ])
        if generated.returncode != 0:
            print(generated.stdout)
            return generated.returncode

        cases = json.loads(cases_path.read_text(encoding="utf-8"))["cases"]
        expected_statuses = Counter(case["expectedStatus"] for case in cases)
        executed = _run([
            sys.executable, "runner/runner.py", str(spec_path), str(SPEC),
            "--base-url", base_url, "--output", str(result_path),
        ])
        outcomes = [
            match.groups()
            for line in executed.stdout.splitlines()
            if (match := re.match(r"\s+\[(OK|NG|WARN)\] (TC-[^ ]+) [A-Z]+ .+ → (\S+)", line))
        ]
        counts = Counter(verdict for verdict, _, _ in outcomes)
        actual_statuses = Counter(status for _, _, status in outcomes)
        print(f"Generated cases: {len(cases)}")
        print(f"Expected status groups: {dict(expected_statuses)}")
        print(f"HTTP results: {dict(counts)} (total {len(outcomes)})")
        print(f"Actual HTTP statuses: {dict(actual_statuses)}")
        if executed.returncode == 0:
            input_errors = _check_combined_required_inputs(cases, result_path)
            if input_errors:
                print("Combined required-input check failed:")
                for error in input_errors:
                    print(f"- {error}")
                return 1
            print("Combined required-input check: passed")
        if executed.returncode != 0:
            for verdict, case_id, _ in outcomes:
                if verdict != "OK":
                    print(f"[{verdict}] {case_id}")
            print("Runner output:")
            print(executed.stdout)
            return executed.returncode
        print(f"Result workbook: {result_path}")
        return 0
    finally:
        api.terminate()
        try:
            api.wait(timeout=5)
        except subprocess.TimeoutExpired:
            api.kill()
            api.wait(timeout=5)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, help="keep generated files in this folder")
    args = parser.parse_args()
    if args.output_dir:
        return _execute(args.output_dir.resolve())
    with tempfile.TemporaryDirectory(prefix="openapi-validation-e2e-") as temporary_dir:
        return _execute(Path(temporary_dir))


if __name__ == "__main__":
    sys.exit(main())
