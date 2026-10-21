"""db_scenario_runner.py の実行結果をExcelレポートとして出力する。

Layer1（openapi-test-run）の結果Excelと違い、Layer2はこれまで標準出力の
ログしか残らなかった。QA記録・チーム内共有として持ち出せる成果物が無かったため、
シナリオ名・判定・APIステータス・詳細を一覧できるシートを新規に作る。
"""
from __future__ import annotations
from datetime import datetime, timezone
from pathlib import Path
import json

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

import sys

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from src.renderer.excel_output import COLOR
from request_evidence import write_request_sheet
from report_layout import safe_cell, style_table, write_overview

HEADERS = ["No", "試験ID", "シナリオ名", "判定", "APIステータス", "詳細"]

_HEADER_FILL = PatternFill(fill_type="solid", fgColor=COLOR["headerBg"])
_HEADER_FONT = Font(color=COLOR["headerFg"], bold=True)
_OK_FILL = PatternFill(fill_type="solid", fgColor=COLOR["normalBg"])
_NG_FILL = PatternFill(fill_type="solid", fgColor=COLOR["abnormalBg"])
_WARNING_FILL = PatternFill(fill_type="solid", fgColor=COLOR["warningBg"] if "warningBg" in COLOR else "FFF2CC")
_THIN_BORDER = Border(
    top=Side(style="thin", color=COLOR["border"]),
    left=Side(style="thin", color=COLOR["border"]),
    bottom=Side(style="thin", color=COLOR["border"]),
    right=Side(style="thin", color=COLOR["border"]),
)
_COL_WIDTHS = [6, 22, 30, 8, 16, 60]


def write_report(rows: list[tuple], output_path: str | Path, evidence: list[dict] | None = None) -> None:
    """Accept legacy 4-field rows and case-ID-aware 5-field rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DBシナリオ結果"

    ws.append([f"実行日時: {datetime.now(timezone.utc).isoformat(timespec='seconds')}"])
    ws.append(HEADERS)
    header_row = ws[2]
    for cell in header_row:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    overview = []
    evidence_by_id = {item.get('case_id'): item for item in evidence or [] if item.get('case_id')}
    for i, row_data in enumerate(rows, start=1):
        if len(row_data) == 5:
            case_id, name, result, api_status, detail = row_data
        elif len(row_data) == 4:
            name, result, api_status, detail = row_data
            case_id = ""
        else:
            raise ValueError("DBシナリオ結果の行は4要素または5要素で指定してください")
        verdict = result if isinstance(result, str) else "OK" if result else "NG"
        row = [safe_cell(value) for value in [i, case_id, name, verdict, api_status, detail]]
        ws.append(row)
        entry = evidence_by_id.get(case_id, {})
        overview.append({'number': i, 'title': name, 'verdict': verdict, 'expected': entry.get('expected_status'),
                         'actual': api_status, 'detail': detail, 'sheet': ws.title, 'row': i + 2})
        fill = _OK_FILL if verdict == "OK" else _WARNING_FILL if verdict == "未確認" else _NG_FILL
        for cell in ws[i + 2]:
            cell.fill = fill
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for col_idx, width in enumerate(_COL_WIDTHS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A3"  # 実行日時行・ヘッダ行を固定し、シナリオ数が多くてもスクロール時に見失わないようにする
    style_table(ws, _COL_WIDTHS, 2, 4)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    if evidence is not None:
        requests = []
        for entry in evidence:
            for phase in ('before', 'main', 'cleanup'):
                if phase == 'main':
                    attempts = entry.get('poll_attempts') or [{'status': entry.get('api_status'), 'request': entry.get('sent_request', {})}]
                    steps = [{'name': '本試験', 'verdict': entry.get('verdict'), 'attempts': attempts}]
                else:
                    steps = [step for step in entry.get('api_steps', []) if step.get('phase', 'before') == phase]
                for step in steps:
                    for index, attempt in enumerate(step.get('attempts', []), 1):
                        requests.append({'case_id': (f"{entry.get('case_id')} / {step['name']} #{index}"
                                                     if entry.get("api_steps") or entry.get("poll_attempts") else entry.get("case_id")),
                                         'verdict': step.get('verdict'), 'status': attempt.get('status'),
                                         'request': attempt.get('request', {})})
        write_request_sheet(wb, requests)
        flow_sheet = wb.create_sheet('API手順と完了待ち')
        flow_sheet.append(['試験ID', '手順', 'API操作', '判定', '期待status', '実測status', '送信回数', '期待レスポンス', '実測レスポンス', '詳細'])
        for entry in evidence:
            steps = list(entry.get('api_steps', []))
            if entry.get('poll_attempts'):
                attempts = entry['poll_attempts']
                steps.append({'name': '本試験の完了待ち', 'verdict': entry.get('verdict'),
                              'attempts': attempts, 'actualStatus': attempts[-1]['status'],
                              'actualBody': attempts[-1]['body'], 'detail': entry.get('detail')})
            for step in steps:
                flow_sheet.append([safe_cell(value) for value in [entry.get('case_id'), step.get('name'), step.get('operationId'),
                    step.get('verdict'), step.get('expectedStatus'), step.get('actualStatus'), len(step.get('attempts', [])),
                    json.dumps(step.get('expectedBody', {}), ensure_ascii=False, default=str),
                    json.dumps(step.get('actualBody', {}), ensure_ascii=False, default=str), step.get('detail')]])
        style_table(flow_sheet, [22, 25, 25, 12, 16, 16, 12, 45, 45, 55], 1, 4)
        comparison = wb.create_sheet('DB比較詳細')
        comparison.append(['期待値と実際の値。長い内容はセルを選択して数式バーで確認できます。記録のない比較は表示しません。'])
        comparison.append(['試験ID', '試験項目', '比較対象', '期待値', '実際の値'])
        for entry in evidence:
            for key, label in [('db_response_comparison', 'DB行とレスポンス'), ('db_insert_comparison', '追加されたDB行'), ('db_sql_response_comparison', 'SQL結果とレスポンス')]:
                data = entry.get(key, {})
                value_keys = ('expected', 'actual') if key == 'db_sql_response_comparison' else ('expected_rows', 'added_rows' if key == 'db_insert_comparison' else 'actual_rows')
                for expected, actual, name in [('expected_count', 'actual_count', '件数'), (*value_keys, '値')]:
                    if expected not in data and actual not in data:
                        continue
                    comparison.append([safe_cell(value) for value in [entry.get('case_id'), entry.get('name'), f'{label}：{name}',
                        json.dumps(data.get(expected, '記録なし'), ensure_ascii=False, indent=2, default=str),
                        json.dumps(data.get(actual, '記録なし'), ensure_ascii=False, indent=2, default=str)]])
        style_table(comparison, [22, 35, 25, 65, 65], 2)
        rules = wb.create_sheet('期待値ルール照合')
        rules.append(['試験ID', '照合先', '期待値', '実際の値', '判定', '詳細'])
        steps = wb.create_sheet('期待値の作成過程')
        steps.append(['試験ID', '処理名', '実行前の参照・計算結果'])
        bulk = wb.create_sheet('InfluxDB一括照合')
        bulk.append(['試験ID', '比較対象', '期待値', '実際の値', '判定'])
        comparison_targets = {}
        dump = lambda value: json.dumps(value, ensure_ascii=False, indent=2, default=str)
        for entry in evidence:
            check = entry.get('expectation_rules', {})
            for step in check.get('steps', []):
                steps.append([safe_cell(v) for v in [entry.get('case_id'), step['name'], dump(step['value'])]])
            for item in check.get('assertions', []):
                definition = item.get('definition', {})
                target = ('レスポンス：' + definition['path']) if 'path' in definition else 'DB：' + str(definition.get('query', {}).get('table', '')) + (' / ' + definition['value_path'] if definition.get('value_path') else ' / 全行')
                detail = item.get('detail', '')
                if 'query' in definition:
                    detail += '\nDB検索条件:\n' + dump(definition['query'])
                rules.append([safe_cell(v) for v in [entry.get('case_id'), target, dump(item.get('expected')),
                    dump(item.get('actual', '記録なし')), item.get('verdict'), detail]])
                previous = comparison_targets.get(entry.get('case_id'))
                if previous is None or (item.get('verdict') == 'NG' and previous[2] != 'NG'):
                    comparison_targets[entry.get('case_id')] = (rules.title, rules.max_row, item.get('verdict'))
            data = entry.get('influx_mutation', {})
            if 'scope' in data.get('definition', {}):
                for label, expected, actual in [
                    ('確認範囲', data['definition']['scope'], '範囲外は未照合'),
                    ('変更対象', data['definition']['target'], {'対象数': data.get('target_count'), '対象外数': data.get('guard_count')}),
                    ('変更前', data.get('before_points'), '期待値を作る基準'),
                    ('変更後の全ポイント', data.get('expected_points'), data.get('after_points', '記録なし'))]:
                    bulk.append([safe_cell(v) for v in [entry.get('case_id'), label, dump(expected), dump(actual), data.get('verdict') if label == '変更後の全ポイント' else '参考']])
                    if label == '変更後の全ポイント':
                        previous = comparison_targets.get(entry.get('case_id'))
                        if previous is None or previous[2] != 'NG':
                            comparison_targets[entry.get('case_id')] = (bulk.title, bulk.max_row, data.get('verdict'))
        style_table(rules, [22, 35, 55, 55, 12, 45], 1, 5)
        style_table(steps, [22, 25, 80], 1)
        style_table(bulk, [22, 25, 60, 60, 12], 1, 5)
        for item, row_data in zip(overview, rows, strict=True):
            case_id = row_data[0] if len(row_data) == 5 else ''
            target = comparison_targets.get(case_id)
            # Keep status/schema failures pointing at the main failure reason.
            if target and (item['verdict'] == 'OK' or target[2] == 'NG'):
                item.update(sheet=target[0], row=target[1])
        for sheet, header in [(flow_sheet, 1), (comparison, 2), (rules, 1), (steps, 1), (bulk, 1)]:
            if sheet.max_row <= header:
                wb.remove(sheet)
    write_overview(wb, overview)
    wb.save(str(output_path))
