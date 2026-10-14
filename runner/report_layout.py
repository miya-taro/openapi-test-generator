"""Shared, readable workbook navigation without changing execution evidence."""
from collections import Counter
from math import ceil
from unicodedata import east_asian_width

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import PageSetupProperties

from sensitive_data import redact_sensitive_text

COLORS = {'OK': 'E8F4EC', 'NG': 'FCE8E6', 'ERROR': 'FCE8E6',
          'WARNING': 'FFF4D6', 'WARN': 'FFF4D6', '未確認': 'FFF4D6', '未実施': 'EEF1F5'}


def safe_cell(value):
    if not isinstance(value, str):
        return value
    value = redact_sensitive_text(value)
    if len(value) > 32000:
        value = value[:32000] + '\n［Excelの文字数上限により省略］'
    return "'" + value if value.startswith(('=', '+', '-', '@', '\t', '\r')) else value


def style_table(sheet, widths, header_row, verdict_column=None):
    sheet.sheet_view.showGridLines = False
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for cell in sheet[header_row]:
        cell.fill = PatternFill('solid', fgColor='243B53')
        cell.font = Font(name='Yu Gothic', color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[header_row].height = 30
    for row in sheet.iter_rows(min_row=header_row + 1):
        verdict = row[verdict_column - 1].value if verdict_column else None
        lines = 1
        for cell, width in zip(row, widths, strict=False):
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.font = Font(name='Yu Gothic', size=10, bold=cell.column == verdict_column)
            cell.fill = PatternFill('solid', fgColor=COLORS.get(verdict, 'FFFFFF' if cell.row % 2 else 'F5F7FA'))
            text = str(cell.value or '')
            lines = max(lines, sum(max(1, ceil(sum(2 if east_asian_width(c) in 'WF' else 1 for c in line) / max(1, width - 2))) for line in text.split('\n')))
        sheet.row_dimensions[row[0].row].height = min(210, max(32, lines * 15 + 8))
    sheet.freeze_panes = f'C{header_row + 1}'
    sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(len(widths))}{max(header_row, sheet.max_row)}'
    sheet.print_title_rows = f'1:{header_row}'
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def write_overview(workbook, rows):
    """One row per planned case, including cases that were not executed."""
    if '結果サマリー' in workbook.sheetnames:
        del workbook['結果サマリー']
    sheet = workbook.create_sheet('結果サマリー', 0)
    counts = Counter(row['verdict'] for row in rows)
    sheet.append(['API試験結果'])
    sheet.merge_cells('A1:H1')
    sheet['A1'].font = Font(name='Yu Gothic', size=18, bold=True, color='243B53')
    sheet.row_dimensions[1].height = 34
    sheet.append([f'全{len(rows)}件　' + ' ／ '.join(f'{label} {count}件' for label, count in counts.items())])
    sheet.merge_cells('A2:H2')
    sheet.append(['要確認の結果を上に表示。未確認・未実施は合格ではありません。「詳細を開く」から元の記録へ進めます。'])
    sheet.merge_cells('A3:H3')
    sheet['A3'].alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[3].height = 32
    sheet.append(['No.', '試験項目', '判定', '期待ステータス', '実際のステータス', '応答時間(ms)', '判定理由・確認箇所', '詳細'])
    order = {'NG': 0, 'ERROR': 0, 'WARNING': 1, 'WARN': 1, '未確認': 1, '未実施': 2, 'OK': 3}
    for row in sorted(rows, key=lambda r: order.get(r['verdict'], 1)):
        sheet.append([safe_cell(value) for value in [row['number'], row['title'], row['verdict'], row.get('expected') or '記録なし',
                      row.get('actual') or '応答なし／未実行', row.get('elapsed'), row.get('detail') or ('設定した照合に合格' if row['verdict'] == 'OK' else '詳細記録を確認してください'), '詳細を開く']])
        link = sheet.cell(sheet.max_row, 8)
        target = row['sheet'].replace("'", "''")
        link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{target}'!A{row['row']}", display='詳細を開く')
    style_table(sheet, [7, 42, 13, 20, 20, 16, 65, 16], 4, 3)
    for row in sheet.iter_rows(min_row=5):
        row[7].font = Font(name='Yu Gothic', color='0563C1', underline='single')
    workbook.active = 0
