"""Readable evidence derived from the actual prepared HTTP request."""
from __future__ import annotations

import shlex
from urllib.parse import quote

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sensitive_data import redact_header_value, redact_sensitive_text

HEADERS = ['試験ID', '判定', 'HTTPステータス', 'メソッド', '送信URL',
           '送信curl（Bash形式）', '送信ヘッダ', '送信ボディ', '記録状態']


def capture_request(prepared, secrets=()) -> dict:
    if prepared is None:
        return {}

    def safe(value):
        text = redact_sensitive_text(str(value))
        for secret in secrets:
            if secret:
                text = text.replace(secret, '[REDACTED]').replace(quote(secret, safe=''), '[REDACTED]')
        return text

    body = prepared.body
    if isinstance(body, bytes):
        try:
            body = body.decode('utf-8')
        except UnicodeDecodeError:
            body = '<バイナリ本文：curl表示では再現できません>'
    return {
        'method': prepared.method,
        'url': safe(prepared.url),
        'headers': {name: safe(redact_header_value(name, value)) for name, value in prepared.headers.items()},
        'body': None if body is None else safe(body),
    }


def readable_curl(request: dict) -> str:
    if not request or not request.get('url'):
        return ''
    parts = ['curl', '--request', request['method'], '--url', shlex.quote(request['url'])]
    for name, value in request.get('headers', {}).items():
        parts.extend(['--header', shlex.quote(f'{name}: {value}')])
    if request.get('body') is not None:
        parts.extend(['--data-binary', shlex.quote(request['body'])])
    return ' '.join(parts)


def write_request_sheet(workbook, rows: list[dict], *, first: bool = False) -> None:
    if '送信リクエスト' in workbook.sheetnames:
        del workbook['送信リクエスト']
    sheet = workbook.create_sheet('送信リクエスト', 0 if first else None)
    sheet.append(['実際のHTTPリクエストから記録。curlはBash形式です。認証情報は伏せ字のため再実行時に入力が必要です。'])
    sheet.append(HEADERS)
    for cell in sheet[2]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243B53')
    for row in rows:
        request = row.get('request') or {}
        values = [row.get('case_id'), row.get('verdict'), row.get('status'), request.get('method'),
                  request.get('url'), readable_curl(request),
                  '\n'.join(f'{key}: {value}' for key, value in request.get('headers', {}).items()),
                  request.get('body'), ('送信を試行（応答なし）' if row.get('status') in (0, 'ERR') else '送信内容を記録') if request else '送信内容なし（未実行・組立失敗など）']
        cells = []
        for value in values:
            if isinstance(value, str):
                if len(value) > 32000:
                    value = value[:32000] + '\n[Excelの文字数制限により省略]'
                if value.startswith(('=', '+', '-', '@', '\t', '\r')):
                    value = "'" + value
            cells.append(value)
        sheet.append(cells)
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        sheet.row_dimensions[sheet.max_row].height = 100
    for index, width in enumerate([22, 10, 16, 12, 55, 100, 55, 65, 30], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A3'
    sheet.auto_filter.ref = f'A2:I{sheet.max_row}'
    if first:
        workbook.active = 0
