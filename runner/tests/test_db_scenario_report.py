import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent.parent))
from db_scenario_report import write_report, HEADERS


def test_write_report_creates_file_with_headers(tmp_path):
    output = tmp_path / "report.xlsx"
    write_report([("シナリオA", True, "201", "")], output)
    assert output.exists()

    wb = openpyxl.load_workbook(str(output))
    ws = wb["DBシナリオ結果"]
    header_values = [c.value for c in ws[2]]
    assert header_values == HEADERS


def test_write_report_row_values_match_input(tmp_path):
    output = tmp_path / "report.xlsx"
    write_report(
        [
            ("createUser", True, "201", ""),
            ("deleteUser", False, "404", "該当する行が見つかりません"),
        ],
        output,
    )
    wb = openpyxl.load_workbook(str(output))
    ws = wb["DBシナリオ結果"]
    row1 = [c.value for c in ws[3]]
    row2 = [c.value for c in ws[4]]
    assert row1[:5] == [1, None, "createUser", "OK", "201"]
    assert not row1[5]  # openpyxlは空文字を書き込むとNoneとして読み戻す
    assert row2 == [2, None, "deleteUser", "NG", "404", "該当する行が見つかりません"]


def test_write_report_ok_and_ng_rows_have_different_fill_colors(tmp_path):
    output = tmp_path / "report.xlsx"
    write_report([("ok_case", True, "200", ""), ("ng_case", False, "500", "err")], output)
    wb = openpyxl.load_workbook(str(output))
    ws = wb["DBシナリオ結果"]
    ok_fill = ws[3][0].fill.fgColor.rgb
    ng_fill = ws[4][0].fill.fgColor.rgb
    assert ok_fill != ng_fill


def test_write_report_empty_rows_still_creates_valid_file(tmp_path):
    output = tmp_path / "report.xlsx"
    write_report([], output)
    wb = openpyxl.load_workbook(str(output))
    ws = wb["DBシナリオ結果"]
    assert [c.value for c in ws[2]] == HEADERS
    assert ws.max_row == 2


def test_write_report_creates_parent_directory(tmp_path):
    output = tmp_path / "nested" / "dir" / "report.xlsx"
    write_report([("a", True, "200", "")], output)
    assert output.exists()


def test_write_report_includes_shared_case_id_for_new_rows(tmp_path):
    output = tmp_path / "report.xlsx"
    write_report([("DB-case-1", "createUser", True, "201", "")], output)

    wb = openpyxl.load_workbook(output, data_only=True)
    ws = wb["DBシナリオ結果"]
    assert "試験ID" in [cell.value for cell in ws[2]]
    row = [cell.value for cell in ws[3]]
    assert row[:5] == [1, "DB-case-1", "createUser", "OK", "201"]


def test_db_comparison_overview_keeps_values_and_unverified_separate(tmp_path):
    target = tmp_path/'comparison.xlsx'
    evidence = [{'case_id':'b', 'name':'値の照合', 'expected_status':200,
                 'db_response_comparison': {'expected_rows':[{'value':10}], 'actual_rows':[{'value':99}],
                                            'expected_count':1, 'actual_count':1}}]
    write_report([('a', '準備不足', '未確認', '未実行', '条件の省略を検出できない'),
                  ('b', '=HYPERLINK("x")', 'NG', 200, '値の不一致')], target, evidence)
    wb = openpyxl.load_workbook(target)
    assert wb.active.title == '結果サマリー'
    assert wb.active['C5'].value == 'NG'
    assert wb.active['D5'].value == 200
    assert wb.active['B5'].data_type == 's'
    assert '未確認 1件' in wb.active['A2'].value
    values = wb['DB比較詳細']
    assert '10' in values['D4'].value and '99' in values['E4'].value
    assert values.freeze_panes == 'C3'
    assert wb['DBシナリオ結果']['C4'].data_type == 's'


def test_unused_tabs_removed_and_summary_links_to_actual_mismatch(tmp_path):
    target = tmp_path / 'rules.xlsx'
    evidence = [{'case_id': 'rule', 'name': '入力の10倍', 'expected_status': 200, 'expectation_rules': {
        'steps': [{'name': 'scaled', 'value': 20}], 'assertions': [
            {'definition': {'path': 'value'}, 'expected': 20, 'actual': 20, 'verdict': 'OK'},
            {'definition': {'query': {'table': 'measurements'}, 'value_path': '0.amount'},
             'expected': 20, 'actual': 2, 'verdict': 'NG'}]}}]
    write_report([('rule', '入力の10倍', 'NG', 200, 'DB保存値が不一致')], target, evidence)
    workbook = openpyxl.load_workbook(target)
    assert 'InfluxDB一括照合' not in workbook.sheetnames
    assert 'API手順と完了待ち' not in workbook.sheetnames
    assert 'DB比較詳細' not in workbook.sheetnames
    assert workbook.active['H5'].hyperlink.location == "'期待値ルール照合'!A3"
    assert workbook['期待値ルール照合']['B3'].value == 'DB：measurements / 0.amount'
    assert workbook['期待値ルール照合']['C3'].value == '20'
    assert workbook['期待値ルール照合']['D3'].value == '2'
    assert 'measurements' in workbook['期待値ルール照合']['F3'].value


def test_status_failure_keeps_link_to_status_reason(tmp_path):
    output = tmp_path / 'status.xlsx'
    write_report([('rule', 'status mismatch', 'NG', 500, 'HTTP statusが不一致')], output, [
        {'case_id': 'rule', 'expectation_rules': {'assertions': [
            {'definition': {'path': 'value'}, 'expected': 1, 'actual': 1, 'verdict': 'OK'}]}}])
    workbook = openpyxl.load_workbook(output)
    assert workbook.active['H5'].hyperlink.location == "'DBシナリオ結果'!A3"
