import json
import sys
from pathlib import Path

import pytest
import requests
import yaml

sys.path.insert(0, str(Path(__file__).parent.parent))
sys.path.insert(0, str(Path(__file__).parent.parent.parent))

import db_scenario_runner as dsr_module

MINIMAL_OPENAPI = {
    "openapi": "3.1.0",
    "paths": {
        "/users": {
            "post": {"operationId": "createUser", "responses": {"201": {"description": "Created"}}}
        },
        "/users/{id}": {
            "get": {"operationId": "getUser", "responses": {"200": {"description": "OK"}}},
            "put": {"operationId": "updateUser", "responses": {"200": {"description": "OK"}}}
        },
    },
}

CONFIG = {
    "postgres": {"host": "h", "port": 5432, "dbname": "d", "user": "u", "password": "p"},
    "influxdb": {"host": "h", "port": 8086, "database": "m"},
}


def _write_files(tmp_path, scenarios_data: dict) -> tuple[Path, Path, Path]:
    scenario_path = tmp_path / "scenarios.yaml"
    scenario_path.write_text(yaml.dump(scenarios_data, allow_unicode=True), encoding="utf-8")
    openapi_path = tmp_path / "openapi.json"
    openapi_path.write_text(json.dumps(MINIMAL_OPENAPI), encoding="utf-8")
    config_path = tmp_path / "config.json"
    config_path.write_text(json.dumps(CONFIG), encoding="utf-8")
    return scenario_path, openapi_path, config_path


def run_main(monkeypatch, argv: list[str]) -> int:
    monkeypatch.setattr(sys, "argv", ["db_scenario_runner.py", *argv])
    return dsr_module.main()


def _patch_common(monkeypatch):
    monkeypatch.setattr(dsr_module, "health_check", lambda base_url, health_path: True)


def _fake_call_api(status: int, response_json: dict):
    """main() のオーケストレーションだけを検証するテスト用に、_call_api を丸ごと
    差し替える（HTTP層自体の検証は test_call_api_* / test_http_client.py で行う）。
    """
    def _call_api(base_url, path, method, path_params, body, session=None, persist_cookies=False):
        return status, response_json, None

    return _call_api


# ---------- _resolve_operation ----------


def test_resolve_operation_finds_path_and_method():
    path, method = dsr_module._resolve_operation(MINIMAL_OPENAPI, "createUser")
    assert path == "/users"
    assert method == "POST"


def test_resolve_operation_unknown_raises():
    with pytest.raises(dsr_module.ScenarioError, match="unknownOp"):
        dsr_module._resolve_operation(MINIMAL_OPENAPI, "unknownOp")


# ---------- _call_api ----------


def _fake_response(status: int, body: str) -> requests.Response:
    resp = requests.Response()
    resp.status_code = status
    resp._content = body.encode("utf-8")
    return resp


def test_call_api_builds_url_with_path_params_and_parses_status(monkeypatch):
    captured = {}

    def fake_request(method, url, **kwargs):
        captured["url"] = url
        return _fake_response(201, '{"id": 5}')

    monkeypatch.setattr(requests, "request", fake_request)
    status, response, error = dsr_module._call_api(
        "http://x", "/users/{id}", "PUT", {"id": 5}, {"name": "updated"}
    )
    assert status == 201
    assert response == {"id": 5}
    assert error is None
    assert captured["url"] == "http://x/users/5"


def test_call_api_uses_explicit_profile_auth_header(monkeypatch):
    captured = {}

    def fake_request(self, method, url, **kwargs):
        captured.update(kwargs)
        return _fake_response(200, "{}")

    monkeypatch.setattr(requests.Session, "request", fake_request)
    session = requests.Session()
    dsr_module._call_api(
        "http://x", "/users", "GET", {}, {}, session=session,
        auth_header=("Authorization", "Bearer read-only-token"),
    )
    assert captured["headers"]["Authorization"] == "Bearer read-only-token"


def test_call_api_omits_body_flag_when_body_empty(monkeypatch):
    captured = {}

    def fake_request(method, url, **kwargs):
        captured["kwargs"] = kwargs
        return _fake_response(204, "")

    monkeypatch.setattr(requests, "request", fake_request)
    status, response, error = dsr_module._call_api("http://x", "/users/1", "DELETE", {}, {})
    assert status == 204
    assert "json" not in captured["kwargs"]


def test_call_api_connection_error(monkeypatch):
    def fake_request(method, url, **kwargs):
        raise requests.exceptions.ConnectionError("接続できません")

    monkeypatch.setattr(requests, "request", fake_request)
    status, response, error = dsr_module._call_api("http://x", "/users", "POST", {}, {"a": 1})
    assert error is not None
    assert status == 0


def test_call_api_request_exception_returns_error(monkeypatch):
    def fake_request(method, url, **kwargs):
        raise requests.exceptions.RequestException("connection refused")

    monkeypatch.setattr(requests, "request", fake_request)
    status, response, error = dsr_module._call_api("http://x", "/users", "POST", {}, {})
    assert error == "connection refused"


def test_call_api_invalid_json_body_falls_back_to_empty_dict(monkeypatch):
    def fake_request(method, url, **kwargs):
        return _fake_response(200, "not json")

    monkeypatch.setattr(requests, "request", fake_request)
    status, response, error = dsr_module._call_api("http://x", "/users", "GET", {}, {})
    assert status == 200
    assert response == {}
    assert error is None


# ---------- main(): 引数検証・存在確認 ----------


def test_missing_scenario_file_returns_3(monkeypatch, tmp_path):
    _, openapi_path, config_path = _write_files(tmp_path, {"scenarios": []})
    code = run_main(
        monkeypatch,
        [str(tmp_path / "nope.yaml"), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 3


def test_missing_openapi_file_returns_3(monkeypatch, tmp_path):
    scenario_path, _, config_path = _write_files(tmp_path, {"scenarios": []})
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(tmp_path / "nope.json"), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 3


def test_invalid_scenario_yaml_returns_3(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, {"not_scenarios": []})
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 3


def test_health_check_failure_returns_3(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, {"scenarios": []})
    monkeypatch.setattr(dsr_module, "health_check", lambda base_url, health_path: False)
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 3


def test_skip_health_check_bypasses_a_failing_health_check(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, {"scenarios": []})
    monkeypatch.setattr(
        dsr_module, "health_check",
        lambda base_url, health_path: (_ for _ in ()).throw(AssertionError("health_check should not be called")),
    )
    code = run_main(
        monkeypatch,
        [
            str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path),
            "--skip-health-check",
        ],
    )
    assert code == 0


def test_zero_scenarios_returns_0(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, {"scenarios": []})
    _patch_common(monkeypatch)
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0


# ---------- main(): シナリオ実行の分岐（postgres_backend/influx_backendをフェイクに差し替え） ----------


def _one_postgres_scenario():
    return {
        "scenarios": [
            {
                "name": "ユーザー作成",
                "request": {"operationId": "createUser", "body": {"name": "太郎"}},
                "db_check": {"engine": "postgres", "table": "users", "where": {"id": "{{response.id}}"}, "expect": {}},
            }
        ]
    }


def test_scenario_pass_returns_0(monkeypatch, tmp_path, capsys):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"id": 5}))
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0
    assert "[OK]" in capsys.readouterr().out


def test_post_cleanup_deletes_only_rows_created_after_the_snapshot(monkeypatch, tmp_path, capsys):
    scenario_data = {"scenarios": [{
        "name": "POST行を片付ける",
        "request": {"operationId": "createUser", "body": {"email": "new@example.com"}},
        "response_check": {"expect": {"status": 201}},
        "db_check": {
            "engine": "postgres", "table": "users", "where": {"email": "new@example.com"},
            "expect": {"exists": True},
        },
        "cleanup": {
            "table": "users", "primary_key": "id", "where": {"email": "new@example.com"},
        },
    }]}
    scenario_path, openapi_path, config_path = _write_files(tmp_path, scenario_data)
    _patch_common(monkeypatch)
    snapshots = iter([[4], [4, 9]])
    monkeypatch.setattr(
        dsr_module.postgres_backend,
        "matching_primary_keys",
        lambda *args: next(snapshots),
    )
    deleted = []
    monkeypatch.setattr(
        dsr_module.db_factory, "delete", lambda config, table, primary_key: deleted.append((table, primary_key))
    )
    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"id": 9}))
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )

    assert code == 0
    assert deleted == [("users", {"id": 9})]
    assert "新規作成されたDB行を1件片付けました" in capsys.readouterr().out


def test_case_ids_runs_only_selected_self_contained_db_scenario(monkeypatch, tmp_path):
    scenario_data = {"scenarios": [
        {"id": "one", "name": "first", "request": {"operationId": "createUser"},
         "response_check": {"expect": {"status": 201}}},
        {"id": "two", "name": "second", "request": {"operationId": "createUser"},
         "response_check": {"expect": {"status": 201}}},
    ]}
    scenario_path, openapi_path, config_path = _write_files(tmp_path, scenario_data)
    _patch_common(monkeypatch)
    calls = []

    def fake_api(*args, **kwargs):
        calls.append(args[2])
        return 201, {}, None

    monkeypatch.setattr(dsr_module, "_call_api", fake_api)
    code = run_main(monkeypatch, [
        str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path),
        "--case-ids", "DB-two",
    ])

    assert code == 0
    assert len(calls) == 1


def test_case_ids_reject_response_dependent_db_scenario(monkeypatch, tmp_path, capsys):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)
    called = []
    monkeypatch.setattr(dsr_module, "_call_api", lambda *args, **kwargs: called.append(True))

    code = run_main(monkeypatch, [
        str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path),
        "--case-ids", dsr_module.load_scenarios(scenario_path)[0].case_id,
    ])

    assert code == 4
    assert not called
    assert "単独再実行できません" in capsys.readouterr().err


def test_case_ids_reject_scenario_with_response_dependent_given(monkeypatch, tmp_path, capsys):
    scenario_data = {"scenarios": [{
        "id": "child-data",
        "name": "前ケースのIDで前提データを作る",
        "given": [{"table": "users", "as": "user", "fields": {"id": "{{response.id}}"}}],
        "request": {"operationId": "createUser"},
        "response_check": {"expect": {"status": 201}},
    }]}
    scenario_path, openapi_path, config_path = _write_files(tmp_path, scenario_data)
    _patch_common(monkeypatch)
    called = []
    monkeypatch.setattr(dsr_module, "_call_api", lambda *args, **kwargs: called.append(True))

    code = run_main(monkeypatch, [
        str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path),
        "--case-ids", "DB-child-data",
    ])

    assert code == 4
    assert not called
    assert "単独再実行できません" in capsys.readouterr().err


def test_report_output_writes_excel_with_one_row_per_scenario(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"id": 5}))
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    report_path = tmp_path / "report.xlsx"
    code = run_main(
        monkeypatch,
        [
            str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path),
            "--report-output", str(report_path),
        ],
    )
    assert code == 0
    assert report_path.exists()

    import openpyxl
    wb = openpyxl.load_workbook(str(report_path))
    ws = wb.active
    assert ws[3][1].value.startswith("DB-")
    assert ws[3][2].value == "ユーザー作成"
    assert ws[3][3].value == "OK"


def test_report_output_not_written_when_flag_omitted(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"id": 5}))
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0
    assert not (tmp_path / "report.xlsx").exists()


def test_junit_xml_flag_writes_file(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"id": 5}))
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    junit_path = tmp_path / "result.xml"
    code = run_main(
        monkeypatch,
        [
            str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path),
            "--junit-xml", str(junit_path),
        ],
    )
    assert code == 0
    assert junit_path.exists()


def test_without_junit_xml_flag_no_file_written(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"id": 5}))
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0
    assert not (tmp_path / "result.xml").exists()


def test_scenario_fail_returns_1_and_prints_detail(monkeypatch, tmp_path, capsys):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"id": 5}))
    monkeypatch.setattr(
        dsr_module.postgres_backend, "check", lambda check, config: (False, "該当する行が見つかりません")
    )

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 1
    out = capsys.readouterr().out
    assert "[NG]" in out
    assert "該当する行が見つかりません" in out


def test_scenario_placeholder_resolution_passes_resolved_where_to_backend(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"id": 42}))
    captured = {}

    def fake_check(check, config):
        captured["where"] = check.where
        return True, ""

    monkeypatch.setattr(dsr_module.postgres_backend, "check", fake_check)

    run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert captured["where"] == {"id": 42}


# ---------- シナリオ間の {{response.xxx}} 連携（path_params / body） ----------


def _create_then_update_scenarios():
    return {
        "scenarios": [
            {
                "name": "ユーザー作成",
                "request": {"operationId": "createUser", "body": {"name": "太郎"}},
                "db_check": {"engine": "postgres", "table": "users", "where": {"id": "{{response.id}}"}, "expect": {}},
            },
            {
                "name": "作成したユーザーを更新",
                "request": {
                    "operationId": "updateUser",
                    "path_params": {"id": "{{response.id}}"},
                    "body": {"name": "次郎"},
                },
                "db_check": {"engine": "postgres", "table": "users", "where": {"id": "{{response.id}}"}, "expect": {}},
            },
        ]
    }


def test_path_params_resolved_from_previous_scenario_response(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _create_then_update_scenarios())
    _patch_common(monkeypatch)
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    calls = []

    def fake_call_api(base_url, path, method, path_params, body, session=None, persist_cookies=False):
        calls.append(path_params)
        if len(calls) == 1:
            return 201, {"id": 42}, None
        return 200, {"id": 42, "name": "次郎"}, None

    monkeypatch.setattr(dsr_module, "_call_api", fake_call_api)

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0
    assert len(calls) == 2
    # path_params の {{response.id}} が前段シナリオの id=42 に解決されていること
    assert calls[1] == {"id": 42}


def test_body_resolved_from_previous_scenario_response(monkeypatch, tmp_path):
    data = {
        "scenarios": [
            {
                "name": "ユーザー作成",
                "request": {"operationId": "createUser", "body": {"name": "太郎"}},
                "db_check": {"engine": "postgres", "table": "users", "where": {}, "expect": {}},
            },
            {
                "name": "作成したユーザーIDを紐づけて注文作成",
                "request": {"operationId": "createUser", "body": {"linked_user_id": "{{response.id}}"}},
                "db_check": {"engine": "postgres", "table": "orders", "where": {}, "expect": {}},
            },
        ]
    }
    scenario_path, openapi_path, config_path = _write_files(tmp_path, data)
    _patch_common(monkeypatch)
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    calls = []

    def fake_call_api(base_url, path, method, path_params, body, session=None, persist_cookies=False):
        calls.append(body)
        return 201, {"id": 99}, None

    monkeypatch.setattr(dsr_module, "_call_api", fake_call_api)

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0
    assert calls[1] == {"linked_user_id": 99}


def test_path_params_placeholder_with_no_previous_scenario_marks_ng(monkeypatch, tmp_path, capsys):
    # 1件目からいきなり {{response.id}} を参照している場合、前段のレスポンスが空なので解決に失敗する
    data = {
        "scenarios": [
            {
                "name": "前段なしでいきなり更新",
                "request": {"operationId": "updateUser", "path_params": {"id": "{{response.id}}"}, "body": {}},
                "db_check": {"engine": "postgres", "table": "users", "where": {}, "expect": {}},
            }
        ]
    }
    scenario_path, openapi_path, config_path = _write_files(tmp_path, data)
    _patch_common(monkeypatch)

    def fail_if_called(base_url, path, method, path_params, body, session=None, persist_cookies=False):
        raise AssertionError("解決に失敗した時点でAPIを呼んではいけない")

    monkeypatch.setattr(dsr_module, "_call_api", fail_if_called)

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 1
    out = capsys.readouterr().out
    assert "[NG]" in out
    assert "request値の解決に失敗" in out


def test_previous_response_updates_to_most_recent_scenario_only(monkeypatch, tmp_path):
    # 3シナリオ連鎖: 2番目の path_params は1番目のレスポンス、3番目の path_params は
    # （1番目ではなく）2番目のレスポンスを参照すること（常に「直前」のレスポンスであること）
    data = {
        "scenarios": [
            {
                "name": "1件目",
                "request": {"operationId": "createUser", "body": {}},
                "db_check": {"engine": "postgres", "table": "users", "where": {}, "expect": {}},
            },
            {
                "name": "2件目",
                "request": {"operationId": "updateUser", "path_params": {"id": "{{response.id}}"}, "body": {}},
                "db_check": {"engine": "postgres", "table": "users", "where": {}, "expect": {}},
            },
            {
                "name": "3件目",
                "request": {"operationId": "updateUser", "path_params": {"id": "{{response.id}}"}, "body": {}},
                "db_check": {"engine": "postgres", "table": "users", "where": {}, "expect": {}},
            },
        ]
    }
    scenario_path, openapi_path, config_path = _write_files(tmp_path, data)
    _patch_common(monkeypatch)
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    calls = []

    def fake_call_api(base_url, path, method, path_params, body, session=None, persist_cookies=False):
        calls.append(path_params)
        # 1件目はid=1を、2件目はid=2を返す。3件目はid=2を参照すべき（id=1ではない）
        responses = [(201, {"id": 1}), (200, {"id": 2}), (200, {"id": 2})]
        status, response_json = responses[len(calls) - 1]
        return status, response_json, None

    monkeypatch.setattr(dsr_module, "_call_api", fake_call_api)

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0
    assert calls[1] == {"id": 1}
    assert calls[2] == {"id": 2}


def test_unresolvable_placeholder_marks_scenario_ng(monkeypatch, tmp_path, capsys):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"unrelated_field": 1}))  # id が返らない
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 1
    assert "[NG]" in capsys.readouterr().out


def test_unknown_operation_id_marks_scenario_ng(monkeypatch, tmp_path):
    data = {
        "scenarios": [
            {
                "name": "存在しないop",
                "request": {"operationId": "doesNotExist", "body": {}},
                "db_check": {"engine": "postgres", "table": "t", "expect": {}},
            }
        ]
    }
    scenario_path, openapi_path, config_path = _write_files(tmp_path, data)
    _patch_common(monkeypatch)
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 1


def test_influxdb_scenario_dispatches_to_influx_backend(monkeypatch, tmp_path):
    data = {
        "scenarios": [
            {
                "name": "センサー",
                "request": {"operationId": "createUser", "body": {}},
                "db_check": {"engine": "influxdb", "measurement": "m", "tags": {}, "expect": {}},
            }
        ]
    }
    scenario_path, openapi_path, config_path = _write_files(tmp_path, data)
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {}))
    called = {"hit": False}

    def fake_influx_check(check, config):
        called["hit"] = True
        return True, ""

    monkeypatch.setattr(dsr_module.influx_backend, "check", fake_influx_check)
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert called.get("hit") is True
    assert code == 0


# ---------- main(): response_check（APIレスポンス自体の検証） ----------


def _one_response_check_scenario():
    return {
        "scenarios": [
            {
                "name": "最大値のユーザーを取得できる",
                "request": {"operationId": "createUser", "body": {}},
                "response_check": {"expect": {"status": 201, "body": {"price": 99999999.99}}},
            }
        ]
    }


def test_response_check_only_scenario_passes(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_response_check_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"price": 99999999.99}))
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0


def test_response_check_only_scenario_fails_with_detail(monkeypatch, tmp_path, capsys):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_response_check_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"price": 1}))  # 期待値(99999999.99)と異なる
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 1
    out = capsys.readouterr().out
    assert "[NG]" in out
    assert "price" in out


def test_response_check_status_mismatch_marks_ng(monkeypatch, tmp_path, capsys):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_response_check_scenario())
    _patch_common(monkeypatch)

    # statusが期待(201)と異なる
    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(404, {"price": 99999999.99}))
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 1
    assert "[NG]" in capsys.readouterr().out


def test_response_check_and_db_check_combined_both_must_pass(monkeypatch, tmp_path, capsys):
    data = {
        "scenarios": [
            {
                "name": "作成もDBもレスポンスも確認",
                "request": {"operationId": "createUser", "body": {}},
                "db_check": {"engine": "postgres", "table": "users", "where": {}, "expect": {}},
                "response_check": {"expect": {"status": 201}},
            }
        ]
    }
    scenario_path, openapi_path, config_path = _write_files(tmp_path, data)
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {}))
    # db_check側はNG、response_check側はOKにする → 全体としてNGになるはず
    monkeypatch.setattr(
        dsr_module.postgres_backend, "check", lambda check, config: (False, "該当する行が見つかりません")
    )

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 1
    out = capsys.readouterr().out
    assert "[NG]" in out
    assert "該当する行が見つかりません" in out


def test_response_check_and_db_check_both_pass_marks_ok(monkeypatch, tmp_path):
    data = {
        "scenarios": [
            {
                "name": "作成もDBもレスポンスも確認",
                "request": {"operationId": "createUser", "body": {}},
                "db_check": {"engine": "postgres", "table": "users", "where": {}, "expect": {}},
                "response_check": {"expect": {"status": 201}},
            }
        ]
    }
    scenario_path, openapi_path, config_path = _write_files(tmp_path, data)
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {}))
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0


# ---------- main(): given（factory前提データ） ----------


def _one_given_scenario():
    return {
        "scenarios": [
            {
                "name": "在庫がある商品を注文すると在庫が減る",
                "given": [
                    {"table": "products", "as": "product_a", "fields": {"name": "商品A", "stock": 10}},
                ],
                "request": {
                    "operationId": "createUser",
                    "body": {"product_id": "{{given.product_a.id}}", "qty": 1},
                },
                "db_check": {
                    "engine": "postgres", "table": "products",
                    "where": {"id": "{{given.product_a.id}}"}, "expect": {"stock": 9},
                },
            }
        ]
    }


def test_given_creates_data_before_api_call_and_cleans_up_after(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_given_scenario())
    _patch_common(monkeypatch)

    calls = []

    def fake_create(pg_config, table, fields):
        calls.append(("create", table, fields))
        return {"id": 5, "stock": 10}, {"id": 5}

    def fake_delete(pg_config, table, primary_key):
        calls.append(("delete", table, primary_key))

    def fake_call_api(base_url, path, method, path_params, body, session=None, persist_cookies=False):
        calls.append(("call_api", path_params, body))
        return 201, {"id": 1}, None

    monkeypatch.setattr(dsr_module.db_factory, "create", fake_create)
    monkeypatch.setattr(dsr_module.db_factory, "delete", fake_delete)
    monkeypatch.setattr(dsr_module, "_call_api", fake_call_api)
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0
    kinds = [c[0] for c in calls]
    assert kinds == ["create", "call_api", "delete"]  # create → API呼び出し → cleanup の順
    assert calls[0][1] == "products"
    assert calls[0][2] == {"name": "商品A", "stock": 10}
    assert calls[2][1] == "products"
    assert calls[2][2] == {"id": 5}


def test_given_data_referenced_in_request_body_and_db_check(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_given_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module.db_factory, "create", lambda pg_config, table, fields: ({"id": 5, "stock": 10}, {"id": 5}))
    monkeypatch.setattr(dsr_module.db_factory, "delete", lambda pg_config, table, primary_key: None)

    captured = {}

    def fake_call_api(base_url, path, method, path_params, body, session=None, persist_cookies=False):
        captured["body"] = body
        return 201, {"id": 1}, None

    monkeypatch.setattr(dsr_module, "_call_api", fake_call_api)

    captured_check = {}

    def fake_check(check, config):
        captured_check["where"] = check.where
        return True, ""

    monkeypatch.setattr(dsr_module.postgres_backend, "check", fake_check)

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0
    assert captured["body"] == {"product_id": 5, "qty": 1}
    assert captured_check["where"] == {"id": 5}


def test_get_response_check_resolves_values_from_given_db_row(monkeypatch, tmp_path):
    data = {
        "scenarios": [{
            "name": "DBにあるユーザーをGETで読み取れる",
            "given": [{
                "table": "users",
                "as": "user",
                "fields": {"name": "Read User", "email": "read@example.com", "role": "member"},
            }],
            "request": {
                "operationId": "getUser",
                "path_params": {"id": "{{given.user.id}}"},
            },
            "response_check": {
                "expect": {
                    "status": 200,
                    "body": {
                        "id": "{{given.user.id}}",
                        "name": "{{given.user.name}}",
                        "email": "{{given.user.email}}",
                    },
                }
            },
        }]
    }
    scenario_path, openapi_path, config_path = _write_files(tmp_path, data)
    _patch_common(monkeypatch)
    monkeypatch.setattr(
        dsr_module.db_factory,
        "create",
        lambda pg_config, table, fields: (
            {"id": 42, **fields},
            {"id": 42},
        ),
    )
    monkeypatch.setattr(dsr_module.db_factory, "delete", lambda pg_config, table, primary_key: None)
    captured = {}

    def fake_call_api(base_url, path, method, path_params, body, session=None, persist_cookies=False):
        captured.update({"path": path, "method": method, "path_params": path_params})
        return 200, {"id": 42, "name": "Read User", "email": "read@example.com", "role": "member"}, None

    monkeypatch.setattr(dsr_module, "_call_api", fake_call_api)

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )

    assert code == 0
    assert captured == {"path": "/users/{id}", "method": "GET", "path_params": {"id": 42}}


def test_given_cleanup_happens_even_on_scenario_failure(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_given_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module.db_factory, "create", lambda pg_config, table, fields: ({"id": 5, "stock": 10}, {"id": 5}))
    deleted = []
    monkeypatch.setattr(dsr_module.db_factory, "delete", lambda pg_config, table, primary_key: deleted.append((table, primary_key)))
    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"id": 1}))
    # db_check側をNGにする
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (False, "stock不一致"))

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 1  # シナリオ自体はNG
    assert deleted == [("products", {"id": 5})]  # それでもcleanupは実行される


def test_given_cleanup_failure_warns_and_returns_warning_exit_code(monkeypatch, tmp_path, capsys):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_given_scenario())
    _patch_common(monkeypatch)

    monkeypatch.setattr(dsr_module.db_factory, "create", lambda pg_config, table, fields: ({"id": 5, "stock": 10}, {"id": 5}))

    def fail_delete(pg_config, table, primary_key):
        raise RuntimeError("connection lost")

    monkeypatch.setattr(dsr_module.db_factory, "delete", fail_delete)
    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {"id": 1}))
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    # シナリオの判定は成功のまま、片付け失敗を示すwarning exit codeを返す
    assert code == 2
    err = capsys.readouterr().err
    assert "[WARN]" in err
    assert "クリーンアップ" in err
    assert "1 件" in err  # 個別[WARN]行に加えて、実行全体のサマリでも件数が可視化される


def test_given_creation_failure_marks_scenario_ng_and_skips_api_call(monkeypatch, tmp_path, capsys):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_given_scenario())
    _patch_common(monkeypatch)

    def fail_create(pg_config, table, fields):
        raise dsr_module.db_factory.FactoryError("NOT NULL制約違反")

    monkeypatch.setattr(dsr_module.db_factory, "create", fail_create)

    def fail_if_called(*a, **k):
        raise AssertionError("前提データ投入に失敗した時点でAPIを呼んではいけない")

    monkeypatch.setattr(dsr_module, "_call_api", fail_if_called)

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 1
    out = capsys.readouterr().out
    assert "[NG]" in out
    assert "前提データの投入に失敗しました" in out


def test_given_multiple_items_cleaned_up_in_reverse_order(monkeypatch, tmp_path):
    data = {
        "scenarios": [
            {
                "name": "カテゴリと商品を両方前提データとして用意する",
                "given": [
                    {"table": "categories", "as": "category_a", "fields": {"name": "カテゴリA"}},
                    {"table": "products", "as": "product_a", "fields": {"category_id": "{{given.category_a.id}}"}},
                ],
                "request": {"operationId": "createUser", "body": {}},
                "db_check": {"engine": "postgres", "table": "products", "where": {}, "expect": {}},
            }
        ]
    }
    scenario_path, openapi_path, config_path = _write_files(tmp_path, data)
    _patch_common(monkeypatch)

    created = []
    def fake_create(pg_config, table, fields):
        created.append((table, fields))
        if table == "categories":
            return {"id": 1, "name": "カテゴリA"}, {"id": 1}
        return {"id": 2, "category_id": fields.get("category_id")}, {"id": 2}

    deleted = []
    monkeypatch.setattr(dsr_module.db_factory, "create", fake_create)
    monkeypatch.setattr(dsr_module.db_factory, "delete", lambda pg_config, table, primary_key: deleted.append((table, primary_key)))
    monkeypatch.setattr(dsr_module, "_call_api", _fake_call_api(201, {}))
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0
    assert created == [
        ("categories", {"name": "カテゴリA"}),
        ("products", {"category_id": 1}),
    ]
    # 作成順(categories → products)の逆順(products → categories)で削除される（FK依存を考慮）
    assert deleted == [("products", {"id": 2}), ("categories", {"id": 1})]


# ---------- main(): --auth-config ----------


def test_auth_config_login_failure_returns_3(monkeypatch, tmp_path):
    # main() は os.environ を直接書き換える。monkeypatch.delenv(raising=False) は
    # キーが元々存在しない場合は何も記録しないため、テスト終了後に main() が新規
    # 設定した値が残ってしまう。setenv で一旦触れておくことで「元は未設定だった」
    # ことを monkeypatch に記録させ、テスト終了時に確実に削除させる。
    monkeypatch.setenv("TEST_AUTH_HEADER_NAME", "__placeholder__")
    monkeypatch.setenv("TEST_AUTH_HEADER_VALUE", "__placeholder__")
    scenario_path, openapi_path, config_path = _write_files(tmp_path, {"scenarios": []})
    _patch_common(monkeypatch)
    auth_config_path = tmp_path / "auth.json"
    auth_config_path.write_text(
        json.dumps({"login_url": "/login", "method": "POST"}), encoding="utf-8"
    )

    def fake_perform_login(base_url, auth_cfg, session=None):
        raise dsr_module.AuthLoginError("ログイン失敗")

    monkeypatch.setattr(dsr_module, "perform_login", fake_perform_login)
    code = run_main(
        monkeypatch,
        [
            str(scenario_path), str(openapi_path), "--base-url", "http://x",
            "--config", str(config_path), "--auth-config", str(auth_config_path),
        ],
    )
    assert code == 3


def test_auth_config_success_sets_env_and_call_api_sends_header(monkeypatch, tmp_path):
    monkeypatch.setenv("TEST_AUTH_HEADER_NAME", "__placeholder__")
    monkeypatch.setenv("TEST_AUTH_HEADER_VALUE", "__placeholder__")
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)
    auth_config_path = tmp_path / "auth.json"
    auth_config_path.write_text(
        json.dumps({"login_url": "/login", "method": "POST"}), encoding="utf-8"
    )

    monkeypatch.setattr(dsr_module, "perform_login", lambda base_url, auth_cfg, session=None: "tok123")
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    # ここでは _call_api 自体は差し替えず実装のまま動かし、実際にHTTP層まで届くヘッダを
    # requests.Session.request のフェイクで捕捉することで、
    # 「ログイン成功→環境変数設定→_call_api がその環境変数を読んでヘッダに載せる」
    # という配線全体を検証する（main() は Session を生成して _call_api に渡すため、
    # モジュールレベルの requests.request ではなく Session.request を差し替える必要がある）。
    captured = {}

    def fake_request(self, method, url, **kwargs):
        captured["headers"] = kwargs.get("headers", {})
        return _fake_response(201, '{"id": 5}')

    monkeypatch.setattr(requests.Session, "request", fake_request)

    code = run_main(
        monkeypatch,
        [
            str(scenario_path), str(openapi_path), "--base-url", "http://x",
            "--config", str(config_path), "--auth-config", str(auth_config_path),
        ],
    )
    assert code == 0
    assert captured["headers"].get("Authorization") == "Bearer tok123"


def test_named_auth_profiles_login_once_and_use_isolated_sessions(monkeypatch, tmp_path, capsys):
    scenario_data = {
        "scenarios": [
            {
                "name": "read-only profile is denied",
                "auth_profile": "read_only",
                "request": {"operationId": "createUser"},
                "response_check": {"expect": {"status": 200}},
            },
            {
                "name": "support profile is denied",
                "auth_profile": "support",
                "request": {"operationId": "createUser"},
                "response_check": {"expect": {"status": 200}},
            },
        ]
    }
    scenario_path, openapi_path, config_path = _write_files(tmp_path, scenario_data)
    _patch_common(monkeypatch)
    auth_config_path = tmp_path / "auth.json"
    auth_config_path.write_text(json.dumps({
        "login_url": "/login",
        "method": "POST",
        "auth_profiles": {
            "read_only": {"username": "reader", "password": "reader-secret"},
            "support": {"username": "support", "password": "support-secret"},
        },
    }), encoding="utf-8")
    monkeypatch.delenv("TEST_AUTH_HEADER_NAME", raising=False)
    monkeypatch.delenv("TEST_AUTH_HEADER_VALUE", raising=False)

    logins = []
    def fake_login(base_url, auth_cfg, session=None):
        logins.append((auth_cfg["body"]["username"], session))
        return f"token-{auth_cfg['body']['username']}"

    captured = []
    def fake_request(self, method, url, **kwargs):
        captured.append((self, kwargs.get("headers", {})))
        return _fake_response(200, "{}")

    monkeypatch.setattr(dsr_module, "perform_login", fake_login)
    monkeypatch.setattr(requests.Session, "request", fake_request)
    code = run_main(monkeypatch, [
        str(scenario_path), str(openapi_path), "--base-url", "http://x",
        "--config", str(config_path), "--auth-config", str(auth_config_path),
    ])

    assert code == 0
    assert [name for name, _ in logins] == ["reader", "support"]
    assert logins[0][1] is not logins[1][1]
    assert captured[0][0] is logins[0][1]
    assert captured[1][0] is logins[1][1]
    assert captured[0][1]["Authorization"] == "Bearer token-reader"
    assert captured[1][1]["Authorization"] == "Bearer token-support"
    output = capsys.readouterr().out
    assert "認証: read_only" in output
    assert "認証: support" in output


def test_auth_profile_without_auth_config_is_rejected_before_http(monkeypatch, tmp_path):
    scenario_data = {"scenarios": [{
        "name": "requires reader login", "auth_profile": "reader",
        "request": {"operationId": "createUser"},
        "response_check": {"expect": {"status": 200}},
    }]}
    scenario_path, openapi_path, config_path = _write_files(tmp_path, scenario_data)
    _patch_common(monkeypatch)
    monkeypatch.setattr(dsr_module, "health_check", lambda *args: pytest.fail("HTTP must not start"))

    code = run_main(monkeypatch, [
        str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path),
    ])

    assert code == 3


# ---------- main(): --cookie-jar ----------


def test_cookie_jar_flag_passes_shared_session_with_persist_cookies_true(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    seen = []

    def fake_call_api(base_url, path, method, path_params, body, session=None, persist_cookies=False):
        seen.append((session, persist_cookies))
        return 201, {"id": 5}, None

    monkeypatch.setattr(dsr_module, "_call_api", fake_call_api)

    code = run_main(
        monkeypatch,
        [
            str(scenario_path), str(openapi_path), "--base-url", "http://x",
            "--config", str(config_path), "--cookie-jar",
        ],
    )
    assert code == 0
    session, persist_cookies = seen[0]
    assert isinstance(session, requests.Session)
    assert persist_cookies is True


def test_without_cookie_jar_flag_session_shared_but_not_persisted(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    _patch_common(monkeypatch)
    monkeypatch.setattr(dsr_module.postgres_backend, "check", lambda check, config: (True, ""))

    seen = []

    def fake_call_api(base_url, path, method, path_params, body, session=None, persist_cookies=False):
        seen.append((session, persist_cookies))
        return 201, {"id": 5}, None

    monkeypatch.setattr(dsr_module, "_call_api", fake_call_api)

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path)],
    )
    assert code == 0
    session, persist_cookies = seen[0]
    assert isinstance(session, requests.Session)  # コネクションプーリングのため常に共有される
    assert persist_cookies is False


# ---------- main(): --dry-run ----------


def test_dry_run_returns_0_without_executing(monkeypatch, tmp_path, capsys):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, _one_postgres_scenario())
    # dry-run はネットワーク・DBに触れないため、触れたら失敗するようフェイクを仕込む
    monkeypatch.setattr(dsr_module, "health_check", lambda base_url, health_path: (_ for _ in ()).throw(AssertionError("called")))
    monkeypatch.setattr(
        dsr_module, "_call_api",
        lambda *a, **k: (_ for _ in ()).throw(AssertionError("called")),
    )
    monkeypatch.setattr(
        dsr_module.postgres_backend, "check",
        lambda check, config: (_ for _ in ()).throw(AssertionError("called")),
    )

    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path), "--dry-run"],
    )
    assert code == 0
    out = capsys.readouterr().out
    assert "シナリオ数: 1 件" in out
    assert "DRY-RUN" in out


def test_dry_run_still_validates_scenario_file(monkeypatch, tmp_path):
    scenario_path, openapi_path, config_path = _write_files(tmp_path, {"not_scenarios": []})
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(config_path), "--dry-run"],
    )
    assert code == 3


def test_dry_run_still_validates_config_file(monkeypatch, tmp_path):
    scenario_path, openapi_path, _ = _write_files(tmp_path, {"scenarios": []})
    bad_config_path = tmp_path / "bad_config.json"
    bad_config_path.write_text("{not valid json", encoding="utf-8")
    code = run_main(
        monkeypatch,
        [str(scenario_path), str(openapi_path), "--base-url", "http://x", "--config", str(bad_config_path), "--dry-run"],
    )
    assert code == 3
