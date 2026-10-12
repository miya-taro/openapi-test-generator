"""Request evidence must describe the sent payload, without shell-escape noise or secrets."""
import json
import shlex
import sys
from pathlib import Path

import openpyxl
import requests

sys.path.insert(0, str(Path(__file__).parents[1]))
from request_evidence import capture_request, readable_curl
from excel_writer import write_results
from models import RunCase, RunResult


def test_curl_round_trips_nested_json_and_quoted_values():
    body = {'nested': {'name': 'O\'Brien "quoted"', 'items': [1, 2]}}
    prepared = requests.Request('POST', 'https://example.test/v1/data?site_id=12', json=body).prepare()
    record = capture_request(prepared)
    tokens = shlex.split(readable_curl(record))
    assert tokens[tokens.index('--url') + 1] == prepared.url
    assert tokens[tokens.index('--data-binary') + 1] == prepared.body.decode()
    assert json.loads(tokens[tokens.index('--data-binary') + 1]) == body
    assert '{BASE_URL}' not in readable_curl(record)


def test_login_request_redacts_body_custom_header_and_url():
    secret = 'test-login-secret'
    prepared = requests.Request('POST', f'https://example.test/login?session={secret}',
                                headers={'custom-auth': secret}, json={'username': 'test-user', 'password': secret}).prepare()
    record = capture_request(prepared, [secret])
    assert secret not in json.dumps(record)
    assert secret not in readable_curl(record)
    assert 'test-user' in readable_curl(record)


def test_result_workbook_replaces_preview_and_shows_actual_curl(tmp_path):
    source = tmp_path / 'spec.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'test_spec'
    sheet.append([]); sheet.append([]); sheet.append([])
    sheet.append(['試験ID', 'curl コマンド', '判定'])
    sheet.append(['case-1', 'curl {BASE_URL} \\"preview\\"', ''])
    sheet.append(['case-2', 'curl {BASE_URL}', ''])
    book.save(source)
    case = RunCase('case-1', 'create', '/v1/data', 'POST', 'body', 'requestBody', '', '200', '', '', '', '')
    prepared = requests.Request('POST', 'https://example.test/v1/data', json={'items': [{'id': 1}]}).prepare()
    record = capture_request(prepared)
    result = RunResult(case, 200, '{}', {}, 1, 0, curl_command='curl.exe Windows command', request_evidence=record)
    dest = tmp_path / 'result.xlsx'
    write_results(source, [(result, 'OK', '')], dest)
    output = openpyxl.load_workbook(dest)
    assert output.active.title == '送信リクエスト'
    sheet = output['test_spec']
    assert sheet.cell(4, 2).value == '送信curl（Bash形式）'
    assert sheet.cell(5, 2).value == readable_curl(record)
    assert sheet.cell(6, 2).value == '未実行'
    all_values = '\n'.join(str(v) for ws in output for row in ws.values for v in row)
    assert '{BASE_URL}' not in all_values
    assert '\\"' not in sheet.cell(5, 2).value
    assert output.active.cell(3, 5).value == prepared.url
    assert json.loads(output.active.cell(3, 8).value) == {'items': [{'id': 1}]}


def test_recorded_curl_replays_the_actual_http_payload(tmp_path, monkeypatch):
    import shutil
    import subprocess
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    import pytest
    from http_client import execute_with_openapi

    executable = shutil.which('curl')
    if not executable:
        pytest.skip('curl required for replay verification')
    for name in ('TEST_AUTH_TOKEN', 'TEST_AUTH_HEADER_NAME', 'TEST_AUTH_HEADER_VALUE'):
        monkeypatch.delenv(name, raising=False)
    captured = []

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            captured.append((self.path, self.rfile.read(int(self.headers['Content-Length']))))
            self.send_response(200)
            self.send_header('Content-Length', '2')
            self.end_headers()
            self.wfile.write(b'{}')

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        payload = {'data': [{'name': 'O\'Brien "quoted"', 'value': 3.5}]}
        case = RunCase('live', 'echo', '/echo', 'POST', 'body', 'requestBody', json.dumps(payload), '200', '', '', '', '')
        api = {'paths': {'/echo': {'post': {'operationId': 'echo', 'requestBody': {'content': {'application/json': {'schema': {'type': 'object'}}}}}}}}
        result = execute_with_openapi(case, f'http://127.0.0.1:{server.server_port}', api)
        assert result.actual_status == 200
        tokens = shlex.split(readable_curl(result.request_evidence))
        tokens[0] = executable
        replay = subprocess.run(tokens, capture_output=True, timeout=10)
        assert replay.returncode == 0
        assert replay.stdout == b'{}'
        assert len(captured) == 2
        assert captured[0] == captured[1]
        assert json.loads(captured[0][1]) == payload
        from db_scenario_runner import _call_api
        from db_scenario_report import write_report
        db_result = _call_api(f'http://127.0.0.1:{server.server_port}', '/echo', 'POST', {}, payload)
        assert db_result.status == 200
        assert captured[2] == captured[0]
        report = tmp_path / 'db-result.xlsx'
        write_report([('db-live', 'echo', 'OK', '200', '')], report,
                     [{'case_id': 'db-live', 'verdict': 'OK', 'api_status': '200',
                       'sent_request': db_result.request_evidence}])
        workbook = openpyxl.load_workbook(report)
        assert workbook['送信リクエスト'].cell(3, 5).value.endswith('/echo')
        assert json.loads(workbook['送信リクエスト'].cell(3, 8).value) == payload
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)
