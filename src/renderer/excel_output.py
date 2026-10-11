"""Excel 出力（test_spec シート）。openpyxl を使用。"""
from __future__ import annotations
import json
from urllib.parse import quote, urlencode

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter

from src.cmd_escape import quote_cmd_arg as _quote_cmd
from src.security_payload import AUTH_NONE_MARKER, OPENAPI_SECURITY_PARAM, parse_security_payload
from src.coverage import write_coverage_sheet

# ── 列定義 ─────────────────────────────────────────────────────────
# (header, key, compute, width) — key と compute はどちらか一方のみ指定する
COLUMNS: list[tuple[str, str | None, object, int]] = [
    ("試験ID", "id", None, 18),
    ("operationId", "operationId", None, 20),
    ("概要", "summary", None, 36),
    ("パス", "path", None, 22),
    ("メソッド", "method", None, 10),
    ("in", "in", None, 12),
    ("パラメータ名", "paramName", None, 20),
    ("Keyword", "keyword", None, 20),
    ("観点", "perspective", None, 24),
    ("入力/操作", "inputValue", None, 30),
    ("期待ステータス", "expectedStatus", None, 14),
    ("期待レスポンスボディ", "expectedResponseBody", None, 36),
    ("期待レスポンスヘッダ", "expectedResponseHeader", None, 30),
    ("期待応答時間", "expectedResponseTime", None, 14),
    ("期待結果", "expectedResult", None, 30),
    ("判定", "verdict", None, 10),
    ("備考", "notes", None, 36),
    ("curl コマンド", None, True, 80),
]

_OMIT_MARKERS = {"（省略）", "（フィールドを省略）", AUTH_NONE_MARKER}
_EMPTY_MARKER = "（空文字）"

# Excel/CSV フォーミュラインジェクション対策: これらの文字で始まるセル値は
# Excel が数式として評価しうるため、先頭にアポストロフィを付与して文字列扱いにする
_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@", "\t", "\r")


def _sanitize_cell_value(value: object) -> object:
    if isinstance(value, str) and value.startswith(_FORMULA_TRIGGER_CHARS):
        return "'" + value
    return value


def _encode_uri_component(value: str) -> str:
    # JS の encodeURIComponent は A-Z a-z 0-9 - _ . ! ~ * ' ( ) をエスケープしない
    return quote(value, safe="!~*'()")


def _build_curl_command(tc: dict) -> str:
    method = tc["method"].upper()
    omit = tc["inputValue"] in _OMIT_MARKERS
    raw_val = "" if tc["inputValue"] == _EMPTY_MARKER else tc["inputValue"]

    url_val = raw_val[1:-1] if raw_val.startswith('"') and raw_val.endswith('"') and len(raw_val) >= 2 else raw_val

    url = "{BASE_URL}" + tc["path"]
    security_payload = None
    if tc.get("paramName") == OPENAPI_SECURITY_PARAM and raw_val != AUTH_NONE_MARKER:
        security_payload = parse_security_payload(raw_val)

    if tc["in"] == "path" and not omit:
        url = url.replace("{" + tc["paramName"] + "}", _encode_uri_component(url_val))

    if tc["in"] == "query" and not omit:
        url += f"?{tc['paramName']}={_encode_uri_component(url_val)}"

    if security_payload and security_payload["query"]:
        query = urlencode(security_payload["query"])
        url += ("&" if "?" in url else "?") + query

    cmd = f'curl.exe -s -v -X {method} {_quote_cmd(url)}'

    if tc["in"] == "header" and not omit and tc.get("paramName") != OPENAPI_SECURITY_PARAM:
        cmd += f' -H {_quote_cmd(f"{tc['paramName']}: {raw_val}")}'

    if security_payload:
        for name, value in security_payload["headers"].items():
            cmd += f' -H {_quote_cmd(f"{name}: {value}")}'
        if security_payload["cookies"]:
            cookie_value = "; ".join(f"{name}={value}" for name, value in security_payload["cookies"].items())
            cmd += f' -H {_quote_cmd(f"Cookie: {cookie_value}")}'

    if tc["in"] == "body" and not omit:
        try:
            body_val = "" if raw_val == "" else json.loads(raw_val)
        except (ValueError, TypeError):
            body_val = raw_val
        body = json.dumps({tc["paramName"]: body_val}, ensure_ascii=False)
        cmd += f' -H "Content-Type: application/json" -d {_quote_cmd(body)}'

    return cmd


# ── 色定数 ─────────────────────────────────────────────────────────
COLOR = {
    "headerBg": "FF1F3864",
    "headerFg": "FFFFFFFF",
    "metaBg": "FFE8EDF4",
    "metaLabel": "FF555555",
    "normalBg": "FFE2EFDA",
    "abnormalBg": "FFFCE4D6",
    "verdictBg": "FFFFF2CC",
    "border": "FFAAAAAA",
}

_THIN_BORDER = Border(
    top=Side(style="thin", color=COLOR["border"]),
    left=Side(style="thin", color=COLOR["border"]),
    bottom=Side(style="thin", color=COLOR["border"]),
    right=Side(style="thin", color=COLOR["border"]),
)


def _write_meta_rows(sheet, meta: dict, col_count: int) -> None:
    sheet.cell(row=1, column=1, value="Source")
    sheet.cell(row=1, column=2, value=meta["source"])
    sheet.cell(row=1, column=4, value="Generated At")
    sheet.cell(row=1, column=5, value=meta["generatedAt"])

    sheet.cell(row=2, column=1, value="Total Cases")
    sheet.cell(row=2, column=2, value=meta["totalCases"])
    sheet.cell(row=2, column=4, value="Normal")
    sheet.cell(row=2, column=5, value=meta["normalCases"])
    sheet.cell(row=2, column=7, value="Abnormal")
    sheet.cell(row=2, column=8, value=meta["abnormalCases"])

    for row_idx in (1, 2):
        for c in range(1, col_count + 1):
            cell = sheet.cell(row=row_idx, column=c)
            cell.fill = PatternFill(fill_type="solid", fgColor=COLOR["metaBg"])
            cell.font = Font(color=COLOR["metaLabel"], size=10)
        for c in (1, 4, 7):
            cell = sheet.cell(row=row_idx, column=c)
            cell.font = Font(bold=True, color=COLOR["metaLabel"], size=10)
        sheet.row_dimensions[row_idx].height = 16

    sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)
    sheet.merge_cells(start_row=1, start_column=5, end_row=1, end_column=6)
    sheet.merge_cells(start_row=2, start_column=2, end_row=2, end_column=3)
    sheet.merge_cells(start_row=2, start_column=5, end_row=2, end_column=6)


def _write_note_row(sheet, col_count: int) -> None:
    note_cell = sheet.cell(row=3, column=1)
    note_cell.value = (
        "※ 本シートはインターフェース観点のテストケースです。"
        "Layer2以降のシナリオテストは別シートを参照してください。"
    )
    note_cell.font = Font(italic=True, color="FF555555", size=9)
    note_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[3].height = 16
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=col_count)


def _write_header_row(sheet) -> None:
    for i, (header, *_rest) in enumerate(COLUMNS):
        cell = sheet.cell(row=4, column=i + 1)
        cell.value = header
        cell.fill = PatternFill(fill_type="solid", fgColor=COLOR["headerBg"])
        cell.font = Font(bold=True, color=COLOR["headerFg"], size=10)
        cell.alignment = Alignment(vertical="center", horizontal="center")
        cell.border = _THIN_BORDER
    sheet.row_dimensions[4].height = 20


def _write_data_rows(sheet, cases: list[dict], verdict_col_idx: int) -> None:
    verdict_values = ["OK", "NG", "未実施"]

    for row_offset, tc in enumerate(cases):
        row_num = row_offset + 5  # 1=meta1, 2=meta2, 3=note, 4=header
        is_normal = tc["perspective"].startswith("正常系")
        row_bg = COLOR["normalBg"] if is_normal else COLOR["abnormalBg"]

        for col_idx, (_header, key, compute, _width) in enumerate(COLUMNS):
            cell = sheet.cell(row=row_num, column=col_idx + 1)
            cell.value = _build_curl_command(tc) if compute else _sanitize_cell_value(tc.get(key) or "")
            cell.font = Font(size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _THIN_BORDER

            is_verdict = (col_idx + 1) == verdict_col_idx
            cell.fill = PatternFill(fill_type="solid", fgColor=COLOR["verdictBg"] if is_verdict else row_bg)

        sheet.row_dimensions[row_num].height = 32

    # 判定列にドロップダウンを設定
    if cases:
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(verdict_values) + '"',
            allow_blank=True,
            error="選択肢から選んでください",
            errorTitle="入力エラー",
            prompt="OK / NG / 未実施 から選択",
        )
        sheet.add_data_validation(dv)
        col_letter = get_column_letter(verdict_col_idx)
        dv.add(f"{col_letter}5:{col_letter}{4 + len(cases)}")


def _configure_sheet(sheet) -> None:
    for i, (_header, _key, _compute, width) in enumerate(COLUMNS):
        sheet.column_dimensions[get_column_letter(i + 1)].width = width
    # 行1〜4（meta×2 + 注意書き + ヘッダ）を固定
    sheet.freeze_panes = "A5"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def write_excel_output(output: dict, excel_path: str) -> None:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "test_spec"

    verdict_col_idx = next(i for i, (_h, key, _c, _w) in enumerate(COLUMNS) if key == "verdict") + 1

    _write_meta_rows(sheet, output["meta"], len(COLUMNS))
    _write_note_row(sheet, len(COLUMNS))
    _write_header_row(sheet)
    _write_data_rows(sheet, output["cases"], verdict_col_idx)
    _configure_sheet(sheet)

    if "coverage" in output:
        write_coverage_sheet(wb, output["coverage"])
    wb.save(excel_path)
