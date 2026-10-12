"""Human-readable API specification sheets; execution IDs stay hidden."""
import json
import shlex

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ['No.', '確認すること', '対象項目', '送信内容', '期待ステータス', '期待応答時間', '判定', '送信予定curl（Bash形式）', '内部ID']


def simple_request(command):
    if not command.startswith('curl '):
        return '生成できません（OpenAPI・Base URLを確認）', command
    args = shlex.split(command)
    url = args[args.index('--url') + 1]
    method = args[args.index('--request') + 1]
    body = args[args.index('--data-binary') + 1] if '--data-binary' in args else None
    parts = ['curl', shlex.quote(url)]
    if method != ('POST' if body is not None else 'GET'):
        parts += ['-X', method]
    for i, arg in enumerate(args):
        if arg == '--header':
            header = args[i + 1]
            name, _, value = header.partition(':')
            if name.lower() in {'content-length', 'user-agent', 'accept-encoding', 'connection'}:
                continue
            if name.lower() == 'content-type' and value.strip().lower().split(';')[0] == 'application/json':
                header = 'Content-Type: application/json'
            parts += ['-H', shlex.quote(header)]
    if body is not None:
        # --data-raw preserves literal @ at the beginning of scalar bodies.
        parts += ['--data-raw' if body.startswith('@') else '-d', shlex.quote(body)]
    return body if body is not None else '本文なし', ' '.join(parts)


def purpose(case):
    target = 'ボディ全体' if case['paramName'] == 'requestBody' else case['paramName']
    if case['paramName'] == '__valid_request__':
        return '設定した正常値でリクエストを送る'
    value = case['inputValue']
    if value in {'（フィールドを省略）', '（省略）'}:
        return f'{target}を送らない場合を確認する'
    if value == 'null':
        return f'{target}にnullを指定した場合を確認する'
    if value == '（空文字）':
        return f'{target}に空文字を指定した場合を確認する'
    names = {'integer': '整数', 'fractional': '小数', 'boolean': '真偽値', 'array': '配列', 'object': 'オブジェクト', 'string': '文字列'}
    if case.get('keyword') == 'type':
        kind = case.get('perspective', '').split('_')[-1]
        return f'{target}に異なる型（{names.get(kind, "文字列")}）を指定する'
    return case.get('summary', '入力条件を確認する').replace('requestBody', 'ボディ全体').replace('additionalProperties', '未定義項目の許可設定')


def schema_lines(schema, path='本文', required=False):
    if not isinstance(schema, dict) or not schema:
        return [f'{path}：制約未設定']
    types = {'object': 'オブジェクト', 'array': '配列', 'string': '文字列', 'integer': '整数', 'number': '数値', 'boolean': '真偽値'}
    lines = [f'{path}：{types.get(schema.get("type"), str(schema.get("type", "型未指定")))}' + ('・必須' if required else '')]
    labels = {'format': '形式', 'minLength': '最小文字数', 'maxLength': '最大文字数', 'minimum': '最小値', 'maximum': '最大値', 'enum': '許可値', 'pattern': '正規表現', 'minItems': '最小要素数', 'maxItems': '最大要素数'}
    for key, label in labels.items():
        if key in schema:
            lines.append(f'{path}の{label}：{json.dumps(schema[key], ensure_ascii=False)}')
    for name, child in schema.get('properties', {}).items():
        lines += schema_lines(child, f'{path}.{name}', name in schema.get('required', []))
    if 'items' in schema:
        lines += schema_lines(schema['items'], path + '[]')
    return lines


def write_readable_sheets(wb, cases, api, base_url, commands):
    groups = {}
    for case, command in zip(cases, commands, strict=True):
        groups.setdefault((case['path'], case['method']), []).append((case, command))
    for index, ((path, method), rows) in enumerate(groups.items(), 1):
        # Stable display-only ordering; retain case/command pairs and execution IDs.
        rows = sorted(rows, key=lambda row: not str(row[0]['expectedStatus']).strip().startswith('2'))
        ws = wb.create_sheet(f'試験仕様_{index:03}', index - 1)
        ws.append(['対象API', f'{method.upper()} {path}'])
        ws.append(['接続先', base_url or '未設定'])
        ws.append(['説明', '送信内容は予定値。認証情報は伏字。判定は実行後に更新。共通条件は「レスポンス条件」シート参照。'])
        ws.append(HEADERS)
        for n, (case, command) in enumerate(rows, 1):
            body, curl = simple_request(command)
            location = {'body': 'ボディ', 'query': 'クエリ', 'header': 'ヘッダ', 'cookie': 'Cookie', 'path': 'パス', 'operation': 'リクエスト全体'}.get(case['in'], case['in'])
            name = case['paramName']
            target = location if name in {'requestBody', '__valid_request__'} else f'{location}：{name}'
            status = str(case['expectedStatus'])
            if status.lower() in {'2xx', '4xx'}:
                responses = (api or {}).get('paths', {}).get(path, {}).get(method.lower(), {}).get('responses', {})
                codes = [str(c) for c in responses if str(c).isdigit() and str(c).startswith(status[0])]
                wildcard = status.upper() in responses or status.lower() in responses or (status.lower() == '4xx' and 'default' in responses)
                status = (status + '（範囲指定）') if wildcard else (' / '.join(sorted(codes)) if codes else status + '（対応するレスポンス定義なし）')
            ws.append([f'{n:03}', purpose(case), target, body, status, case.get('expectedResponseTime', ''), '未実施', curl, case['id']])
        ws.column_dimensions['I'].hidden = True
        style(ws, [8, 46, 25, 50, 30, 18, 12, 85, 20])
    conditions = wb.create_sheet('レスポンス条件')
    conditions.append(['対象API', 'ステータス', 'レスポンスヘッダ', 'レスポンスボディ'])
    for path, method in groups:
        responses = (api or {}).get('paths', {}).get(path, {}).get(method.lower(), {}).get('responses', {})
        if not responses:
            conditions.append([f'{method} {path}', '未設定', '未設定・検証対象外', '未設定・検証対象外'])
        for status, response in responses.items():
            headers = []
            for name, header in response.get('headers', {}).items():
                headers += schema_lines(header.get('schema', {}), name)
            bodies = []
            for media, content in response.get('content', {}).items():
                headers.append(f'Content-Type：{media}')
                bodies += schema_lines(content.get('schema', {}))
            conditions.append([f'{method.upper()} {path}', str(status), '\n'.join(headers) or '未設定・検証対象外', '\n'.join(bodies) or ('本文なし' if not response.get('content') else '未設定・検証対象外')])
    style(conditions, [38, 15, 50, 80], header=1)
    wb.active = 0


def style(ws, widths, header=4):
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                # Explicit strings also protect formulas while keeping JSON readable.
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in ws[header]:
        cell.fill = PatternFill('solid', fgColor='1F3864')
        cell.font = Font(color='FFFFFF', bold=True)
    ws.freeze_panes = f'A{header + 1}'
    ws.auto_filter.ref = f'A{header}:H{ws.max_row}' if header == 4 else f'A1:D{ws.max_row}'
    for n in range(header + 1, ws.max_row + 1):
        ws.row_dimensions[n].height = 75
