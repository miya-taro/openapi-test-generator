import json
import shlex

import openpyxl

from src.renderer.readable_spec import schema_lines, simple_request, write_readable_sheets


def test_openapi31_nullable_and_union_types_in_response_conditions():
    lines = schema_lines({"type": "object", "required": ["value"], "properties": {
        "value": {"type": ["string", "null"]}, "id": {"type": ["integer", "string"]},
    }})
    assert "本文.value：文字列 / null・必須" in lines
    assert "本文.id：整数 / 文字列" in lines


def test_readable_sheet_preserves_json_and_hides_internal_ids():
    wb = openpyxl.Workbook()
    tc = {'id': 'TC-login-019', 'method': 'POST', 'path': '/login', 'in': 'operation',
          'paramName': '__valid_request__', 'inputValue': '', 'expectedStatus': '200'}
    command = "curl --request POST --url http://example.test/login --header 'Content-Length: 15' --header 'Content-Type: application/json; charset=utf-8' --data-binary '{\"user\":\"test\"}'"
    api = {'paths': {'/login': {'post': {'responses': {'200': {'content': {'application/json': {'schema': {'type': 'object', 'required': ['result'], 'properties': {'result': {'type': 'string'}}}}}}}}}}}
    write_readable_sheets(wb, [tc], api, 'http://example.test', [command])
    ws = wb['試験仕様_001']
    assert ws.cell(5, 1).value == '001'
    assert ws.cell(5, 9).value == tc['id']
    assert ws.column_dimensions['I'].hidden
    assert json.loads(ws.cell(5, 4).value) == {'user': 'test'}
    curl = ws.cell(5, 8).value
    assert 'Content-Length' not in curl and 'charset' not in curl and '--data-binary' not in curl
    args = shlex.split(curl)
    assert json.loads(args[args.index('-d') + 1]) == {'user': 'test'}
    assert '本文.result：文字列・必須' in wb['レスポンス条件'].cell(2, 4).value


def test_scalar_at_literal_does_not_read_local_file():
    body, command = simple_request("curl --request POST --url http://example.test --data-binary '@secret'")
    assert body == '@secret'
    assert '--data-raw' in command
