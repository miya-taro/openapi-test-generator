import json

import openpyxl

from src.renderer.excel_output import (
    write_excel_output,
    _build_curl_command,
    COLUMNS,
    COLOR,
)
from src.renderer.json_output import build_output


def _tc(**overrides) -> dict:
    base = {
        "id": "TC-op-001",
        "operationId": "op",
        "summary": "summary",
        "path": "/users/{id}",
        "method": "GET",
        "in": "path",
        "paramName": "id",
        "keyword": "minimum",
        "perspective": "正常系_最小値",
        "inputValue": "1",
        "expectedStatus": "200",
        "expectedResponseBody": "",
        "expectedResponseHeader": "",
        "expectedResponseTime": "3秒以内",
        "expectedResult": "正常に処理されること",
        "verdict": "",
        "notes": "",
    }
    base.update(overrides)
    return base


def test_preview_uses_runner_body_without_network(monkeypatch):
    import shlex
    import requests
    def no_network(*args, **kwargs):
        raise AssertionError("preview must not send requests")
    monkeypatch.setattr(requests.Session, "send", no_network)
    schema = {"type": "object", "required": ["username", "password"], "properties": {
        "username": {"type": "string", "example": "test"},
        "password": {"type": "string", "example": "pass"}}}
    api = {"paths": {"/v1/auth/login": {"post": {"operationId": "login", "requestBody": {
        "required": True, "content": {"application/json": {"schema": schema}}}}}}}
    tc = _tc(operationId="login", path="/v1/auth/login", method="POST",
             **{"in": "body"}, paramName="requestBody",
             inputValue='{ ...validBody, "__extra__": "value" }')
    for name, value, expected in [
        ("requestBody", tc["inputValue"], {"username", "password", "__extra__"}),
        ("username", "（フィールドを省略）", {"password"}),
    ]:
        cmd = _build_curl_command(dict(tc, paramName=name, inputValue=value), api, "http://example.test:8100")
        args = shlex.split(cmd)
        body = json.loads(args[args.index("--data-binary") + 1])
        assert set(body) == expected
        assert "validBody" not in cmd and "BASE_URL" not in cmd
        assert body["password"] != "pass"
        assert args[args.index("--url") + 1] == "http://example.test:8100/v1/auth/login"
    baseline = dict(tc, **{"in": "operation"}, paramName="__valid_request__", inputValue="（最小有効リクエスト）")
    args = shlex.split(_build_curl_command(baseline, api, "http://example.test:8100"))
    body = json.loads(args[args.index("--data-binary") + 1])
    assert body["username"] == "test"
    assert set(body) == {"username", "password"}


def test_preview_without_context_is_explicit():
    assert "OpenAPIとBase URL" in _build_curl_command(_tc())


# --- write_excel_output ---


def test_write_excel_output_creates_sheet_with_expected_headers(tmp_path):
    output = build_output([_tc()], "openapi.json")
    out_path = tmp_path / "out.xlsx"
    write_excel_output(output, str(out_path))

    wb = openpyxl.load_workbook(str(out_path))
    assert wb.sheetnames == ["test_spec"]
    ws = wb["test_spec"]
    header_row = [c.value for c in ws[4]]
    assert header_row == [col[0] for col in COLUMNS]


def test_write_excel_output_data_row_values(tmp_path):
    output = build_output([_tc(summary="サマリー")], "openapi.json")
    out_path = tmp_path / "out.xlsx"
    write_excel_output(output, str(out_path))

    ws = openpyxl.load_workbook(str(out_path))["test_spec"]
    row = [c.value for c in ws[5]]
    assert row[0] == "TC-op-001"
    assert row[2] == "サマリー"
    assert "OpenAPIとBase URL" in row[-1]


def test_write_excel_output_normal_and_abnormal_row_colors(tmp_path):
    cases = [
        _tc(id="TC-op-001", perspective="正常系_最小値"),
        _tc(id="TC-op-002", perspective="異常系_必須欠落"),
    ]
    output = build_output(cases, "openapi.json")
    out_path = tmp_path / "out.xlsx"
    write_excel_output(output, str(out_path))

    ws = openpyxl.load_workbook(str(out_path))["test_spec"]
    normal_fill = ws.cell(row=5, column=1).fill.fgColor.rgb
    abnormal_fill = ws.cell(row=6, column=1).fill.fgColor.rgb
    assert normal_fill == COLOR["normalBg"]
    assert abnormal_fill == COLOR["abnormalBg"]


def test_write_excel_output_verdict_dropdown_added_when_cases_exist(tmp_path):
    output = build_output([_tc()], "openapi.json")
    out_path = tmp_path / "out.xlsx"
    write_excel_output(output, str(out_path))

    ws = openpyxl.load_workbook(str(out_path))["test_spec"]
    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 1
    assert "OK" in dvs[0].formula1


def test_write_excel_output_no_dropdown_when_no_cases(tmp_path):
    output = build_output([], "openapi.json")
    out_path = tmp_path / "out.xlsx"
    write_excel_output(output, str(out_path))

    ws = openpyxl.load_workbook(str(out_path))["test_spec"]
    assert len(ws.data_validations.dataValidation) == 0


def test_write_excel_output_freeze_panes_and_meta_rows(tmp_path):
    output = build_output([_tc()], "openapi.json")
    out_path = tmp_path / "out.xlsx"
    write_excel_output(output, str(out_path))

    ws = openpyxl.load_workbook(str(out_path))["test_spec"]
    assert ws.freeze_panes == "A5"
    assert ws.cell(row=1, column=1).value == "Source"
    assert ws.cell(row=2, column=2).value == output["meta"]["totalCases"]
