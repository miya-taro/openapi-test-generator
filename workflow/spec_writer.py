"""Write a generic, reviewable specification workbook from the case master."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from src.coverage import write_coverage_sheet
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter

HEADERS = ["試験ID", "層", "試験項目", "API", "入力・試験観点", "期待結果", "出典", "レビュー状態"]
_WIDTHS = [22, 16, 36, 32, 52, 48, 52, 16]


def _json_text(value: object) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _safe_cell(value: object) -> object:
    if not isinstance(value, str):
        return value
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _row(case: dict) -> list[object]:
    definition = case.get("definition") or {}
    request = definition.get("request") or {}
    method = request.get("method", "")
    path = request.get("path", "")
    api = " ".join(str(part) for part in (method, path) if part)
    input_detail = {
        key: value
        for key, value in definition.items()
        if key != "request"
    }
    if request.get("location") or request.get("name") or request.get("value") is not None:
        input_detail["parameter"] = {
            key: request[key]
            for key in ("location", "name", "value")
            if key in request
        }
    request_body = {
        key: request[key]
        for key in ("pathParams", "body")
        if key in request
    }
    if request_body:
        input_detail["request"] = request_body
    source = case.get("source") or {}
    source_text = " / ".join(
        str(source[key]) for key in ("type", "path", "operationId", "scenarioId")
        if source.get(key)
    )
    return [
        case.get("caseId", ""),
        "Layer 1" if case.get("kind") == "contract" else "Layer 2（DB）",
        case.get("title", ""),
        api,
        _json_text(input_detail),
        _json_text(case.get("expected")),
        source_text,
        case.get("reviewStatus", "draft"),
    ]


def write_case_spec_excel(master: dict, output_path: str | Path) -> None:
    """Create a portable base workbook; company templates can map these columns later."""
    if master.get("format") != "api-tester.case-master" or not isinstance(master.get("cases"), list):
        raise ValueError("試験ケースマスターの形式が不正です")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "試験仕様"
    sheet.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="243B53")
    border = Border(*( [Side(style="thin", color="CDD5DF")] * 4 ))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for case in master["cases"]:
        sheet.append([_safe_cell(value) for value in _row(case)])
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 28

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    write_coverage_sheet(workbook, master.get("coverage", []))
    workbook.save(path)
