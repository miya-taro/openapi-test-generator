import json
import hashlib

import pytest
from openpyxl import load_workbook

from workflow.case_master import (
    CaseMasterError,
    apply_case_review_registry,
    build_case_master,
    load_case_master,
    save_case_master_reviews,
    write_case_master,
)
from workflow.spec_writer import write_case_spec_excel


def _write_openapi(path):
    path.write_text(json.dumps({
        "openapi": "3.0.3",
        "info": {"title": "fixture", "version": "1"},
        "paths": {
            "/users": {
                "get": {"operationId": "listUsers", "responses": {"200": {"description": "ok"}}}
            },
            "/users/{id}": {
                "get": {"operationId": "getUser", "responses": {"200": {"description": "ok"}}}
            },
        },
    }), encoding="utf-8")


def _contract_output(case_ids=("TC-listUsers-001",)):
    return {
        "meta": {"totalCases": len(case_ids)},
        "cases": [
            {
                "id": case_id,
                "operationId": "listUsers",
                "summary": "正常系_代表値",
                "path": "/users",
                "method": "GET",
                "in": "query",
                "paramName": "limit",
                "inputValue": "20",
                "expectedStatus": "200",
                "expectedResponseBody": "",
                "expectedResponseHeader": "",
                "expectedResponseTime": "3000ms",
                "expectedResult": "正常",
                "perspective": "正常系",
                "keyword": "type",
                "notes": "",
            }
            for case_id in case_ids
        ],
    }


def _write_scenarios(path, text=None):
    path.write_text(text or """
scenarios:
  defaults:
    engine: postgres
  list:
    - id: user-read
      name: ユーザーを取得できる
      request:
        operationId: getUser
        path_params:
          id: 1
      response_check:
        expect:
          status: 200
          body:
            id: 1
""", encoding="utf-8")


def test_build_master_preserves_contract_case_id_and_expected_fields(tmp_path):
    openapi = tmp_path / "openapi.json"
    _write_openapi(openapi)

    master = build_case_master("sample", _contract_output(), None, openapi)

    assert master["format"] == "api-tester.case-master"
    assert master["schemaVersion"] == 1
    assert master["cases"][0]["caseId"] == "TC-listUsers-001"
    assert master["cases"][0]["kind"] == "contract"
    assert master["cases"][0]["reviewStatus"] == "draft"
    assert master["cases"][0]["definition"]["request"] == {
        "method": "GET",
        "path": "/users",
        "location": "query",
        "name": "limit",
        "value": "20",
    }
    assert master["cases"][0]["expected"]["status"] == "200"


def test_build_master_maps_db_scenario_and_uses_explicit_id(tmp_path):
    openapi = tmp_path / "openapi.json"
    scenarios = tmp_path / "scenarios.yaml"
    _write_openapi(openapi)
    _write_scenarios(scenarios)

    master = build_case_master("sample", None, scenarios, openapi)

    case = master["cases"][0]
    assert case["caseId"] == "DB-user-read"
    assert case["kind"] == "db-scenario"
    assert case["source"]["operationId"] == "getUser"
    assert case["definition"]["request"]["path"] == "/users/{id}"
    assert case["definition"]["request"]["pathParams"] == {"id": 1}
    assert case["expected"]["response"]["status"] == 200


def test_build_master_preserves_get_array_contains_expectation(tmp_path):
    openapi = tmp_path / "openapi.json"
    scenarios = tmp_path / "scenarios.yaml"
    _write_openapi(openapi)
    _write_scenarios(scenarios, """
scenarios:
  - id: users-list
    name: 一覧からDBのユーザーを取得できる
    request: {operationId: listUsers}
    response_check:
      expect:
        status: 200
        contains:
          - path: "$"
            fields:
              id: "{{given.user.id}}"
""")

    master = build_case_master("sample", None, scenarios, openapi)

    case = master["cases"][0]
    assert case["expected"]["response"]["contains"] == [
        {"path": "$", "fields": {"id": "{{given.user.id}}"}}
    ]


def test_named_auth_profile_is_reviewable_but_credentials_are_not_in_case_master(tmp_path):
    openapi = tmp_path / "openapi.json"
    scenarios = tmp_path / "scenarios.yaml"
    _write_openapi(openapi)
    _write_scenarios(scenarios, """
scenarios:
  - id: owner-denied
    name: 閲覧者は他ユーザーの注文を変更できない
    auth_profile: read_only
    request: {operationId: getUser, path_params: {id: 4}}
    response_check: {expect: {status: 403}}
""")

    case = build_case_master("sample", None, scenarios, openapi)["cases"][0]
    encoded = json.dumps(case, ensure_ascii=False)

    assert case["definition"]["authProfile"] == "read_only"
    assert "password" not in encoded.lower()
    assert "secret" not in encoded.lower()


def test_request_response_value_rules_are_visible_in_review_and_excel(tmp_path):
    openapi = tmp_path / "openapi.json"
    scenarios = tmp_path / "scenarios.yaml"
    _write_openapi(openapi)
    _write_scenarios(scenarios, """
scenarios:
  - name: 応答IDが送信IDに一致する
    request: {operationId: getUser, path_params: {id: 4}}
    response_check:
      expect:
        status: 200
        values:
          - {path: id, source: path, parameter: id}
""")
    master = build_case_master("sample", None, scenarios, openapi)
    rule = {"path": "id", "source": "path", "parameter": "id"}
    assert master["cases"][0]["expected"]["response"]["values"] == [rule]
    assert master["cases"][0]["definition"]["responseCheck"]["expectedValues"] == [rule]
    output = tmp_path / "cases.xlsx"
    write_case_spec_excel(master, output)
    workbook = load_workbook(output)
    values = [str(cell.value) for row in workbook.active for cell in row]
    assert any('"parameter": "id"' in value for value in values)


def test_runner_and_master_share_the_same_db_case_id(tmp_path):
    import sys
    from pathlib import Path as LocalPath

    runner_dir = str(LocalPath(__file__).resolve().parents[2] / "runner")
    sys.path.insert(0, runner_dir)
    try:
        from db_scenario import load_scenarios
    finally:
        sys.path.remove(runner_dir)

    openapi = tmp_path / "openapi.json"
    scenarios = tmp_path / "scenarios.yaml"
    _write_openapi(openapi)
    _write_scenarios(scenarios, """
scenarios:
  - name: read user
    request: {operationId: getUser, path_params: {id: 4}}
    response_check: {expect: {status: 200}}
""")

    runtime_case_id = load_scenarios(scenarios)[0].case_id
    master_case_id = build_case_master("sample", None, scenarios, openapi)["cases"][0]["caseId"]

    assert runtime_case_id == master_case_id


def test_build_master_combines_contract_and_db_cases_without_id_collision(tmp_path):
    openapi = tmp_path / "openapi.json"
    scenarios = tmp_path / "scenarios.yaml"
    _write_openapi(openapi)
    _write_scenarios(scenarios)

    master = build_case_master("sample", _contract_output(), scenarios, openapi)

    assert [case["caseId"] for case in master["cases"]] == [
        "TC-listUsers-001", "DB-user-read"
    ]
    assert {case["kind"] for case in master["cases"]} == {"contract", "db-scenario"}


def test_fallback_db_ids_are_stable_when_yaml_key_order_changes(tmp_path):
    openapi = tmp_path / "openapi.json"
    first = tmp_path / "first.yaml"
    second = tmp_path / "second.yaml"
    _write_openapi(openapi)
    _write_scenarios(first, """
scenarios:
  - name: fetch user
    request:
      operationId: getUser
      path_params: {id: 1}
    response_check:
      expect: {status: 200, body: {id: 1}}
""")
    _write_scenarios(second, """
scenarios:
  - response_check:
      expect: {body: {id: 1}, status: 200}
    request:
      path_params: {id: 1}
      operationId: getUser
    name: fetch user
""")

    first_case = build_case_master("sample", None, first, openapi)["cases"][0]
    second_case = build_case_master("sample", None, second, openapi)["cases"][0]

    assert first_case["caseId"] == second_case["caseId"]
    assert first_case["caseId"].startswith("DB-")


def test_duplicate_case_ids_are_rejected(tmp_path):
    openapi = tmp_path / "openapi.json"
    _write_openapi(openapi)

    with pytest.raises(CaseMasterError, match="重複"):
        build_case_master("sample", _contract_output(("TC-1", "TC-1")), None, openapi)


def test_duplicate_explicit_db_ids_are_rejected(tmp_path):
    openapi = tmp_path / "openapi.json"
    scenarios = tmp_path / "scenarios.yaml"
    _write_openapi(openapi)
    _write_scenarios(scenarios, """
scenarios:
  - id: same
    name: first
    request: {operationId: getUser, path_params: {id: 1}}
    response_check: {expect: {status: 200}}
  - id: same
    name: second
    request: {operationId: getUser, path_params: {id: 2}}
    response_check: {expect: {status: 200}}
""")

    with pytest.raises(CaseMasterError, match="重複"):
        build_case_master("sample", None, scenarios, openapi)


def test_unknown_scenario_operation_is_rejected_with_scenario_name(tmp_path):
    openapi = tmp_path / "openapi.json"
    scenarios = tmp_path / "scenarios.yaml"
    _write_openapi(openapi)
    _write_scenarios(scenarios, """
scenarios:
  - name: missing route
    request: {operationId: absentOperation}
    response_check: {expect: {status: 200}}
""")

    with pytest.raises(CaseMasterError, match="missing route"):
        build_case_master("sample", None, scenarios, openapi)


def test_master_never_contains_runtime_configuration(tmp_path):
    openapi = tmp_path / "openapi.json"
    scenarios = tmp_path / "scenarios.yaml"
    _write_openapi(openapi)
    _write_scenarios(scenarios, """
scenarios:
  - id: read-user
    name: read user
    runtime_config:
      app_db_password: do-not-copy-this
      auth_token: do-not-copy-this-either
    request: {operationId: getUser, path_params: {id: 1}}
    response_check: {expect: {status: 200}}
""")

    master = build_case_master("sample", _contract_output(), scenarios, openapi)
    serialized = json.dumps(master).lower()

    assert "runtimeconfig" not in serialized
    assert "authconfig" not in serialized
    assert "app_db_password" not in serialized
    assert "do-not-copy-this" not in serialized


def test_case_review_transitions_are_saved_without_changing_definitions(tmp_path):
    openapi = tmp_path / "openapi.json"
    _write_openapi(openapi)
    master = build_case_master("sample", _contract_output(), None, openapi)
    path = tmp_path / "case-master.json"
    write_case_master(master, path)

    save_case_master_reviews(path, {"TC-listUsers-001": "reviewed"})
    assert load_case_master(path)["cases"][0]["reviewStatus"] == "reviewed"
    save_case_master_reviews(path, {"TC-listUsers-001": "approved"})
    saved_case = load_case_master(path)["cases"][0]

    assert saved_case["reviewStatus"] == "approved"
    assert saved_case["definition"] == master["cases"][0]["definition"]


def test_review_transition_updates_spec_workbook_and_its_integrity_hash(tmp_path):
    openapi = tmp_path / "openapi.json"
    _write_openapi(openapi)
    master = build_case_master("sample", _contract_output(), None, openapi)
    path = tmp_path / "case-master.json"
    workbook_path = tmp_path / "case-spec.xlsx"
    write_case_master(master, path)
    write_case_spec_excel(master, workbook_path)
    master["executionArtifacts"] = {
        "caseSpecExcelSha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
    }
    write_case_master(master, path)

    save_case_master_reviews(path, {"TC-listUsers-001": "reviewed"})

    saved = load_case_master(path)
    sheet = load_workbook(workbook_path, data_only=True).active
    assert sheet.cell(row=2, column=8).value == "reviewed"
    assert saved["executionArtifacts"]["caseSpecExcelSha256"] == hashlib.sha256(
        workbook_path.read_bytes()
    ).hexdigest()


def test_case_review_cannot_skip_reviewed_status_or_update_unknown_ids(tmp_path):
    openapi = tmp_path / "openapi.json"
    _write_openapi(openapi)
    path = tmp_path / "case-master.json"
    write_case_master(build_case_master("sample", _contract_output(), None, openapi), path)

    with pytest.raises(CaseMasterError, match="変更できません"):
        save_case_master_reviews(path, {"TC-listUsers-001": "approved"})
    with pytest.raises(CaseMasterError, match="見つかりません"):
        save_case_master_reviews(path, {"TC-unknown": "reviewed"})


def test_case_master_rejects_definition_edits_after_generation(tmp_path):
    openapi = tmp_path / "openapi.json"
    _write_openapi(openapi)
    path = tmp_path / "case-master.json"
    write_case_master(build_case_master("sample", _contract_output(), None, openapi), path)
    data = json.loads(path.read_text(encoding="utf-8"))
    data["cases"][0]["expected"]["status"] = "500"
    path.write_text(json.dumps(data), encoding="utf-8")

    with pytest.raises(CaseMasterError, match="変更"):
        load_case_master(path)


def test_approved_review_is_reused_only_for_identical_case_definition(tmp_path):
    openapi = tmp_path / "openapi.json"
    _write_openapi(openapi)
    path = tmp_path / "20260923-100000-deadbeef" / "case-master.json"
    write_case_master(build_case_master("sample", _contract_output(), None, openapi), path)
    save_case_master_reviews(path, {"TC-listUsers-001": "reviewed"})
    save_case_master_reviews(path, {"TC-listUsers-001": "approved"})

    same = build_case_master("sample", _contract_output(), None, openapi)
    changed = build_case_master("sample", _contract_output(("TC-listUsers-002",)), None, openapi)
    apply_case_review_registry(same, tmp_path / "case-reviews.json")
    apply_case_review_registry(changed, tmp_path / "case-reviews.json")

    assert same["cases"][0]["reviewStatus"] == "approved"
    assert changed["cases"][0]["reviewStatus"] == "draft"


def test_openapi_source_change_invalidates_case_approval(tmp_path):
    openapi = tmp_path / "openapi.json"
    _write_openapi(openapi)
    first = build_case_master("sample", _contract_output(), None, openapi)
    original_hash = first["cases"][0]["definitionHash"]
    openapi.write_text(openapi.read_text(encoding="utf-8").replace('"description": "ok"', '"description": "changed"'), encoding="utf-8")
    second = build_case_master("sample", _contract_output(), None, openapi)

    assert second["cases"][0]["definitionHash"] != original_hash


def test_no_change_scope_is_visible_in_master_and_excel(tmp_path):
    openapi = tmp_path / "openapi.json"
    _write_openapi(openapi)
    scenarios = tmp_path / "unchanged.yaml"
    scopes = [{"table": "users", "columns": ["id", "name"], "where": {"id": 1}}]
    scenarios.write_text(json.dumps({"scenarios": [{"name": "readonly", "request": {"operationId": "getUser", "path_params": {"id": 1}},
                                                    "response_check": {"expect": {"status": 200}}, "db_unchanged": scopes}]}), encoding="utf-8")
    master = build_case_master("sample", None, scenarios, openapi)
    assert master["cases"][0]["expected"]["databaseUnchanged"] == scopes
    assert master["cases"][0]["definition"]["dbUnchanged"] == scopes
    excel = tmp_path / "case-spec.xlsx"
    write_case_spec_excel(master, excel)
    workbook = load_workbook(excel)
    values = [str(cell.value) for sheet in workbook for row in sheet for cell in row]
    assert any("databaseUnchanged" in value and "users" in value for value in values)


def test_influx_mutation_is_visible_in_review_and_spec(tmp_path):
    import yaml
    spec = tmp_path / 'api.json'
    spec.write_text(json.dumps({'openapi': '3.1.0', 'paths': {'/m': {'post': {'operationId': 'write', 'responses': {'200': {'description': 'ok'}}}}}}))
    check = {'action': 'insert', 'measurement': 'measure', 'time': '2026-01-01T00:00:00Z', 'tags': {'item': 'a'}, 'fields': {'value': 2}}
    source = tmp_path / 'db.yaml'
    source.write_text(yaml.safe_dump({'scenarios': {'list': [{'name': 'write', 'request': {'operationId': 'write'}, 'influx_mutation': check}]}}))
    master = build_case_master('influx', None, source, spec)
    case = master['cases'][0]
    assert case['definition']['influxMutation'] == check
    assert case['expected']['influxMutation'] == check
