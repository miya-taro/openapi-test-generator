import openpyxl

from workflow.spec_writer import HEADERS, write_case_spec_excel


def _master():
    return {
        "format": "api-tester.case-master",
        "schemaVersion": 1,
        "project": {"name": "sample"},
        "cases": [
            {
                "caseId": "TC-users-001",
                "kind": "contract",
                "title": "正常系",
                "source": {"operationId": "listUsers", "path": "openapi.json"},
                "definition": {"request": {"method": "GET", "path": "/users", "value": "20"}},
                "expected": {"status": "200", "result": "正常"},
                "reviewStatus": "draft",
            },
            {
                "caseId": "DB-read-user",
                "kind": "db-scenario",
                "title": "ユーザーを取得",
                "source": {"operationId": "getUser", "path": "scenarios.yaml"},
                "definition": {"request": {"method": "GET", "path": "/users/{id}", "pathParams": {"id": 1}}},
                "expected": {"response": {"status": 200}},
                "reviewStatus": "approved",
            },
        ],
    }


def test_spec_workbook_contains_both_layers_and_review_state(tmp_path):
    path = tmp_path / "spec.xlsx"
    write_case_spec_excel(_master(), path)

    ws = openpyxl.load_workbook(path, data_only=True).active
    assert [cell.value for cell in ws[1]] == HEADERS
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows[0][0:4] == ("TC-users-001", "Layer 1", "正常系", "GET /users")
    assert rows[0][-1] == "draft"
    assert rows[1][0:2] == ("DB-read-user", "Layer 2（DB）")
    assert rows[1][-1] == "approved"


def test_spec_workbook_escapes_formula_like_user_content(tmp_path):
    master = _master()
    master["cases"][0]["title"] = "=HYPERLINK(\"https://bad\",\"click\")"
    path = tmp_path / "spec.xlsx"
    write_case_spec_excel(master, path)

    ws = openpyxl.load_workbook(path, data_only=False).active
    assert ws[2][2].value.startswith("'")
    assert ws[2][2].data_type != "f"
